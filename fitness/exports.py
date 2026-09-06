"""Excel and PDF monthly reports for the accountant."""
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from django.http import HttpResponse
from django.utils.http import content_disposition_header
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

CURRENCY = 'MAD'
COMPANY = 'AUMB'


def _num(value):
    return float(Decimal(str(value or 0)))


def _money_text(value):
    return f'{_num(value):,.2f} {CURRENCY}'


def _filename(year, month, extension):
    return f'AUMB-monthly-report-{year}-{month:02d}.{extension}'


def _file_response(payload, content_type, filename):
    response = HttpResponse(payload, content_type=content_type)
    response['Content-Disposition'] = content_disposition_header(True, filename)
    return response


def _thin_border():
    side = Side(style='thin', color='D6DDD8')
    return Border(left=side, right=side, top=side, bottom=side)


def _apply_sheet_layout(sheet, widths):
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = 'A2'
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.print_title_rows = '1:1'


def _write_header_row(sheet, values):
    header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
    fill = PatternFill('solid', fgColor='172126')
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(1, column, value)
        cell.font = header_font
        cell.fill = fill
        cell.alignment = Alignment(horizontal='left', vertical='center')
        cell.border = _thin_border()
    sheet.row_dimensions[1].height = 22


def _write_money_cell(cell, value, bold=False):
    cell.value = _num(value)
    cell.number_format = '#,##0.00'
    cell.font = Font(name='Calibri', bold=bold, size=11)
    cell.alignment = Alignment(horizontal='right')
    cell.border = _thin_border()


def _write_text_cell(cell, value, bold=False):
    cell.value = '' if value is None else str(value)
    cell.font = Font(name='Calibri', bold=bold, size=11)
    cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    cell.border = _thin_border()


def build_monthly_xlsx(overview, income, trainers):
    workbook = Workbook()
    summary = workbook.active
    summary.title = 'Summary'

    title_font = Font(name='Calibri', bold=True, size=16, color='172126')
    label_font = Font(name='Calibri', size=11, color='4A6057')
    value_font = Font(name='Calibri', bold=True, size=11)
    note_font = Font(name='Calibri', italic=True, size=10, color='6B7C78')
    header_fill = PatternFill('solid', fgColor='172126')
    header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
    net = _num(overview['net'])
    net_fill = PatternFill('solid', fgColor='E0F0E3' if net >= 0 else 'F9E2DE')

    summary['A1'] = COMPANY
    summary['A1'].font = title_font
    summary['A2'] = 'Monthly report for the accountant'
    summary['A2'].font = note_font
    summary.merge_cells('A1:C1')
    summary.merge_cells('A2:C2')

    summary['A4'] = 'Period'
    summary['B4'] = overview['label']
    summary['A5'] = 'Currency'
    summary['B5'] = CURRENCY
    summary['A4'].font = summary['A5'].font = label_font
    summary['B4'].font = summary['B5'].font = value_font

    for column, heading in enumerate(('Line', 'Amount (MAD)', 'Note'), start=1):
        cell = summary.cell(7, column, heading)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = _thin_border()

    lines = [
        ('Cash collected', overview['collected'], 'Payments received this month. This is income.'),
        ('Expected from classes', overview['expected'], 'Class price x current members. Not cash in hand.'),
        ('Outstanding', overview['outstanding'], 'Still owed by members. Not counted as income yet.'),
        ('Operating expenses (bills)', overview['operating_total'], 'Electricity, water, rent, cleaning, supplies...'),
        ('Trainer pay due', overview['trainer_due'], 'Monthly pay due this month.'),
        ('Trainer pay already given', overview['trainer_paid'], 'Amount already marked as paid.'),
        ('Total spend', overview['total_spend'], 'Operating expenses + trainer pay due.'),
        ('Net', overview['net'], 'Cash collected minus total spend.'),
    ]
    for index, (label, amount, note) in enumerate(lines, start=8):
        _write_text_cell(summary.cell(index, 1), label, bold=True)
        _write_money_cell(summary.cell(index, 2), amount, bold=label in {'Cash collected', 'Total spend', 'Net'})
        _write_text_cell(summary.cell(index, 3), note)
        if label == 'Net':
            for column in range(1, 4):
                summary.cell(index, column).fill = net_fill

    summary['A17'] = 'Net is cash collected minus bills and trainer pay due. Outstanding member balances are not income.'
    summary['A17'].font = note_font
    summary.merge_cells('A17:C17')
    _apply_sheet_layout(summary, (32, 18, 62))
    summary.freeze_panes = 'A8'
    summary.print_title_rows = '1:7'
    summary.page_setup.orientation = 'portrait'

    income_sheet = workbook.create_sheet('Income')
    _write_header_row(income_sheet, (
        'Class', 'Type', 'Members', 'Price / member', 'Expected', 'Collected', 'Outstanding',
    ))
    for row_index, item in enumerate(income.get('classes') or [], start=2):
        _write_text_cell(income_sheet.cell(row_index, 1), item['name'], bold=True)
        _write_text_cell(income_sheet.cell(row_index, 2), item.get('class_type_label') or '')
        cell = income_sheet.cell(row_index, 3, item.get('member_count') or 0)
        cell.font = Font(name='Calibri', size=11)
        cell.alignment = Alignment(horizontal='right')
        cell.border = _thin_border()
        _write_money_cell(income_sheet.cell(row_index, 4), item.get('price_per_member'))
        _write_money_cell(income_sheet.cell(row_index, 5), item.get('expected_monthly'))
        _write_money_cell(income_sheet.cell(row_index, 6), item.get('collected'))
        _write_money_cell(income_sheet.cell(row_index, 7), item.get('outstanding'))
    total_row = income_sheet.max_row + 1
    _write_text_cell(income_sheet.cell(total_row, 1), 'Total', bold=True)
    _write_text_cell(income_sheet.cell(total_row, 2), '')
    income_sheet.cell(total_row, 3).border = _thin_border()
    income_sheet.cell(total_row, 4).border = _thin_border()
    _write_money_cell(income_sheet.cell(total_row, 5), income.get('total_expected'), bold=True)
    _write_money_cell(income_sheet.cell(total_row, 6), income.get('total_collected'), bold=True)
    _write_money_cell(income_sheet.cell(total_row, 7), income.get('total_outstanding'), bold=True)
    _apply_sheet_layout(income_sheet, (28, 16, 12, 16, 16, 16, 16))

    bills = workbook.create_sheet('Bills')
    _write_header_row(bills, ('Category', 'Description', 'Amount (MAD)', 'Notes'))
    expenses = overview.get('expenses') or []
    if expenses:
        for row_index, item in enumerate(expenses, start=2):
            _write_text_cell(bills.cell(row_index, 1), item.get('category_label') or item.get('category'))
            _write_text_cell(bills.cell(row_index, 2), item.get('title'))
            _write_money_cell(bills.cell(row_index, 3), item.get('amount'))
            _write_text_cell(bills.cell(row_index, 4), item.get('notes'))
        total_row = bills.max_row + 1
        _write_text_cell(bills.cell(total_row, 1), 'Total', bold=True)
        _write_text_cell(bills.cell(total_row, 2), '')
        _write_money_cell(bills.cell(total_row, 3), overview.get('operating_total'), bold=True)
        _write_text_cell(bills.cell(total_row, 4), '')
    else:
        _write_text_cell(bills.cell(2, 1), 'No operating expenses this month.')
        bills.merge_cells('A2:D2')
    _apply_sheet_layout(bills, (18, 32, 16, 40))

    pay = workbook.create_sheet('Trainer pay')
    _write_header_row(pay, ('Trainer', 'Specialization', 'Usual monthly pay', 'Amount due', 'Paid this month'))
    for row_index, item in enumerate(trainers.get('trainers') or [], start=2):
        _write_text_cell(pay.cell(row_index, 1), item.get('name'), bold=True)
        _write_text_cell(pay.cell(row_index, 2), item.get('specialization'))
        _write_money_cell(pay.cell(row_index, 3), item.get('monthly_pay'))
        _write_money_cell(pay.cell(row_index, 4), item.get('pay_amount'))
        _write_text_cell(pay.cell(row_index, 5), 'Yes' if item.get('is_paid') else 'No')
    total_row = pay.max_row + 1
    _write_text_cell(pay.cell(total_row, 1), 'Total', bold=True)
    _write_text_cell(pay.cell(total_row, 2), '')
    pay.cell(total_row, 3).border = _thin_border()
    _write_money_cell(pay.cell(total_row, 4), trainers.get('total_due'), bold=True)
    _write_text_cell(
        pay.cell(total_row, 5),
        f"{trainers.get('paid_count') or 0} paid / {trainers.get('unpaid_count') or 0} unpaid",
        bold=True,
    )
    _apply_sheet_layout(pay, (24, 22, 20, 16, 22))

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pdf_fonts():
    regular = 'Helvetica'
    bold = 'Helvetica-Bold'
    candidates = (
        (Path(r'C:\Windows\Fonts\arial.ttf'), Path(r'C:\Windows\Fonts\arialbd.ttf')),
        (Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'), Path('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf')),
        (Path('/System/Library/Fonts/Supplemental/Arial.ttf'), Path('/System/Library/Fonts/Supplemental/Arial Bold.ttf')),
    )
    for regular_path, bold_path in candidates:
        if not regular_path.exists():
            continue
        if 'ReportFont' not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont('ReportFont', str(regular_path)))
            if bold_path.exists():
                pdfmetrics.registerFont(TTFont('ReportFont-Bold', str(bold_path)))
            else:
                pdfmetrics.registerFont(TTFont('ReportFont-Bold', str(regular_path)))
        return 'ReportFont', 'ReportFont-Bold'
    return regular, bold


def _table_style(header=True, emphasize_last=False):
    ink = colors.HexColor('#172126')
    line = colors.HexColor('#D6DDD8')
    paper = colors.HexColor('#FAFBF8')
    commands = [
        ('FONTNAME', (0, 0), (-1, 0), 'ReportFont-Bold' if 'ReportFont-Bold' in pdfmetrics.getRegisteredFontNames() else 'Helvetica-Bold'),
        ('FONTNAME', (0, 1), (-1, -1), 'ReportFont' if 'ReportFont' in pdfmetrics.getRegisteredFontNames() else 'Helvetica'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('BACKGROUND', (0, 0), (-1, 0), ink),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('TEXTCOLOR', (0, 1), (-1, -1), ink),
        ('ALIGN', (0, 0), (-1, 0), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.4, line),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, paper]),
    ]
    if not header:
        commands = [item for item in commands if item[0] not in {'BACKGROUND', 'TEXTCOLOR'} or item[2][1] != 0]
    if emphasize_last:
        commands.append(('FONTNAME', (0, -1), (-1, -1), 'ReportFont-Bold' if 'ReportFont-Bold' in pdfmetrics.getRegisteredFontNames() else 'Helvetica-Bold'))
        commands.append(('BACKGROUND', (0, -1), (-1, -1), paper))
    return TableStyle(commands)


def build_monthly_pdf(overview, income, trainers):
    regular, bold = _pdf_fonts()
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=f"{COMPANY} monthly report {overview['label']}",
        author=COMPANY,
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        'ReportTitle',
        parent=styles['Heading1'],
        fontName=bold,
        fontSize=16,
        textColor=colors.HexColor('#172126'),
        spaceAfter=2 * mm,
    )
    subtitle = ParagraphStyle(
        'ReportSubtitle',
        parent=styles['Normal'],
        fontName=regular,
        fontSize=10,
        textColor=colors.HexColor('#4A6057'),
        spaceAfter=6 * mm,
    )
    heading = ParagraphStyle(
        'ReportHeading',
        parent=styles['Heading2'],
        fontName=bold,
        fontSize=11,
        textColor=colors.HexColor('#172126'),
        spaceBefore=6 * mm,
        spaceAfter=3 * mm,
    )
    body = ParagraphStyle(
        'ReportBody',
        parent=styles['Normal'],
        fontName=regular,
        fontSize=8.5,
        textColor=colors.HexColor('#4A6057'),
        leading=12,
        spaceAfter=3 * mm,
    )
    cell = ParagraphStyle(
        'ReportCell',
        parent=styles['Normal'],
        fontName=regular,
        fontSize=8,
        textColor=colors.HexColor('#172126'),
        leading=11,
    )

    def text(value):
        return Paragraph(str(value or '').replace('&', '&amp;'), cell)

    net = _num(overview['net'])
    story = [
        Paragraph(COMPANY, title),
        Paragraph(f"Monthly report for the accountant — {overview['label']} ({CURRENCY})", subtitle),
        Paragraph(
            'Net is cash collected minus operating bills and trainer pay due. '
            'Outstanding member balances are not counted as income yet.',
            body,
        ),
        Paragraph('Profit and loss', heading),
    ]

    summary_rows = [
        [text('Line'), text('Amount'), text('Note')],
        [text('Cash collected'), text(_money_text(overview['collected'])), text('Payments received this month')],
        [text('Expected from classes'), text(_money_text(overview['expected'])), text('Price x current members')],
        [text('Outstanding'), text(_money_text(overview['outstanding'])), text('Still owed by members')],
        [text('Operating expenses (bills)'), text(_money_text(overview['operating_total'])), text('Bills, cleaning, supplies')],
        [text('Trainer pay due'), text(_money_text(overview['trainer_due'])), text('Monthly pay due')],
        [text('Trainer pay already given'), text(_money_text(overview['trainer_paid'])), text('Already marked as paid')],
        [text('Total spend'), text(_money_text(overview['total_spend'])), text('Operating + trainer pay')],
        [text('Net leftover' if net >= 0 else 'Net shortfall'), text(_money_text(overview['net'])), text('Collected minus total spend')],
    ]
    summary_table = Table(summary_rows, colWidths=[55 * mm, 35 * mm, 80 * mm])
    summary_table.setStyle(_table_style(emphasize_last=True))
    story.extend([summary_table, Spacer(1, 4 * mm), Paragraph('Income by class', heading)])

    class_rows = [[text('Class'), text('Members'), text('Expected'), text('Collected'), text('Outstanding')]]
    for item in income.get('classes') or []:
        class_rows.append([
            text(item.get('name')),
            text(item.get('member_count')),
            text(_money_text(item.get('expected_monthly'))),
            text(_money_text(item.get('collected'))),
            text(_money_text(item.get('outstanding'))),
        ])
    if len(class_rows) == 1:
        class_rows.append([text('No classes this month.'), text(''), text(''), text(''), text('')])
    else:
        class_rows.append([
            text('Total'),
            text(''),
            text(_money_text(income.get('total_expected'))),
            text(_money_text(income.get('total_collected'))),
            text(_money_text(income.get('total_outstanding'))),
        ])
    class_table = Table(class_rows, colWidths=[50 * mm, 22 * mm, 36 * mm, 36 * mm, 36 * mm])
    class_table.setStyle(_table_style(emphasize_last=len(class_rows) > 2))
    story.extend([class_table, Spacer(1, 4 * mm), Paragraph('Bills and extra charges', heading)])

    bill_rows = [[text('Category'), text('Description'), text('Amount'), text('Notes')]]
    expenses = overview.get('expenses') or []
    if expenses:
        for item in expenses:
            bill_rows.append([
                text(item.get('category_label') or item.get('category')),
                text(item.get('title')),
                text(_money_text(item.get('amount'))),
                text(item.get('notes')),
            ])
        bill_rows.append([text('Total'), text(''), text(_money_text(overview.get('operating_total'))), text('')])
    else:
        bill_rows.append([text('No operating expenses this month.'), text(''), text(''), text('')])
    bill_table = Table(bill_rows, colWidths=[38 * mm, 52 * mm, 32 * mm, 48 * mm])
    bill_table.setStyle(_table_style(emphasize_last=bool(expenses)))
    story.extend([bill_table, Spacer(1, 4 * mm), Paragraph('Trainer pay', heading)])

    pay_rows = [[text('Trainer'), text('Specialization'), text('Amount due'), text('Paid')]]
    for item in trainers.get('trainers') or []:
        pay_rows.append([
            text(item.get('name')),
            text(item.get('specialization')),
            text(_money_text(item.get('pay_amount'))),
            text('Yes' if item.get('is_paid') else 'No'),
        ])
    if len(pay_rows) == 1:
        pay_rows.append([text('No trainers this month.'), text(''), text(''), text('')])
    else:
        pay_rows.append([
            text('Total'),
            text(f"{trainers.get('paid_count') or 0} paid / {trainers.get('unpaid_count') or 0} unpaid"),
            text(_money_text(trainers.get('total_due'))),
            text(_money_text(trainers.get('total_paid'))),
        ])
    pay_table = Table(pay_rows, colWidths=[50 * mm, 50 * mm, 40 * mm, 30 * mm])
    pay_table.setStyle(_table_style(emphasize_last=len(pay_rows) > 2))
    story.append(pay_table)

    document.build(story)
    return buffer.getvalue()


def monthly_xlsx_response(overview, income, trainers):
    payload = build_monthly_xlsx(overview, income, trainers)
    return _file_response(
        payload,
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        _filename(overview['year'], overview['month'], 'xlsx'),
    )


def monthly_pdf_response(overview, income, trainers):
    payload = build_monthly_pdf(overview, income, trainers)
    return _file_response(payload, 'application/pdf', _filename(overview['year'], overview['month'], 'pdf'))


def _cash_filename(year, month, extension):
    return f'AUMB-cash-log-{year}-{month:02d}.{extension}'


def build_cash_log_xlsx(year, month, label, rows, total):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Cash log'

    title_font = Font(name='Calibri', bold=True, size=16, color='FF172126')
    note_font = Font(name='Calibri', italic=True, size=10, color='FF6B7C78')
    header_font = Font(name='Calibri', bold=True, color='FF172126', size=11)
    header_fill = PatternFill(fill_type='solid', fgColor='FFE4EBE7')
    sheet['A1'] = COMPANY
    sheet['A1'].font = title_font
    sheet['A2'] = f'Cash desk log — {label}'
    sheet['A2'].font = note_font
    sheet.merge_cells('A1:H1')
    sheet.merge_cells('A2:H2')

    headers = ('Receipt', 'Date', 'Member', 'CIN', 'Amount (MAD)', 'Method', 'Received by', 'Notes')
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(4, column, value)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='left', vertical='center')
        cell.border = _thin_border()
    sheet.row_dimensions[4].height = 22

    for index, row in enumerate(rows, start=5):
        _write_text_cell(sheet.cell(index, 1), row.get('receipt_number'))
        _write_text_cell(sheet.cell(index, 2), row.get('received_at'))
        _write_text_cell(sheet.cell(index, 3), row.get('member_name'))
        _write_text_cell(sheet.cell(index, 4), row.get('id_number'))
        _write_money_cell(sheet.cell(index, 5), row.get('amount'))
        _write_text_cell(sheet.cell(index, 6), row.get('payment_method') or 'Cash')
        _write_text_cell(sheet.cell(index, 7), row.get('received_by'))
        _write_text_cell(sheet.cell(index, 8), row.get('notes'))

    total_row = 5 + len(rows)
    _write_text_cell(sheet.cell(total_row, 1), 'Total', bold=True)
    _write_money_cell(sheet.cell(total_row, 5), total, bold=True)
    _apply_sheet_layout(sheet, (14, 22, 28, 16, 16, 12, 18, 32))
    sheet.freeze_panes = 'A5'
    sheet.print_title_rows = '4:4'

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_cash_log_pdf(year, month, label, rows, total):
    regular, bold = _pdf_fonts()
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=f'{COMPANY} cash log {label}',
        author=COMPANY,
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        'CashTitle',
        parent=styles['Heading1'],
        fontName=bold,
        fontSize=16,
        textColor=colors.HexColor('#172126'),
        spaceAfter=2 * mm,
    )
    subtitle = ParagraphStyle(
        'CashSubtitle',
        parent=styles['Normal'],
        fontName=regular,
        fontSize=10,
        textColor=colors.HexColor('#4A6057'),
        spaceAfter=6 * mm,
    )
    cell = ParagraphStyle(
        'CashCell',
        parent=styles['Normal'],
        fontName=regular,
        fontSize=8,
        textColor=colors.HexColor('#172126'),
        leading=11,
    )
    header_cell = ParagraphStyle(
        'CashHeader',
        parent=styles['Normal'],
        fontName=bold,
        fontSize=8,
        textColor=colors.white,
        leading=11,
    )

    def text(value):
        return Paragraph(str(value or '').replace('&', '&amp;'), cell)

    def header(value):
        return Paragraph(str(value or '').replace('&', '&amp;'), header_cell)

    table_rows = [[
        header('Receipt'),
        header('Date'),
        header('Member'),
        header('Amount'),
        header('Received by'),
        header('Notes'),
    ]]
    for row in rows:
        table_rows.append([
            text(row.get('receipt_number')),
            text(row.get('received_at')),
            text(row.get('member_name')),
            text(_money_text(row.get('amount'))),
            text(row.get('received_by')),
            text(row.get('notes')),
        ])
    if rows:
        table_rows.append([
            text('Total'),
            text(f'{len(rows)} payments'),
            text(''),
            text(_money_text(total)),
            text(''),
            text(''),
        ])
    else:
        table_rows.append([text('No cash payments this month.'), text(''), text(''), text(''), text(''), text('')])

    table = Table(table_rows, colWidths=[28 * mm, 32 * mm, 42 * mm, 28 * mm, 28 * mm, 32 * mm])
    table.setStyle(_table_style(emphasize_last=True))
    document.build([
        Paragraph(COMPANY, title),
        Paragraph(f'Cash desk log — {label} ({CURRENCY})', subtitle),
        table,
    ])
    return buffer.getvalue()


def cash_log_xlsx_response(year, month, label, rows, total):
    return _file_response(
        build_cash_log_xlsx(year, month, label, rows, total),
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        _cash_filename(year, month, 'xlsx'),
    )


def cash_log_pdf_response(year, month, label, rows, total):
    return _file_response(
        build_cash_log_pdf(year, month, label, rows, total),
        'application/pdf',
        _cash_filename(year, month, 'pdf'),
    )
