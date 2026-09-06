import json
from datetime import time, timedelta
from decimal import Decimal

from django.contrib.auth.models import Group, User
from django.test import TestCase
from django.utils import timezone

from fitness.models import Attendance, ClassMember, ClassSchedule, FitnessClassType, GymExpense, GymNotification, GymPayment, GymWhatsAppReminder, Membership, MembershipPlan, Trainer, TrainerPayroll, TrainingClass
from fitness.controllers import create_expiring_membership_notifications
from users.models import ClientProfile


class TrainingClassTest(TestCase):
    def setUp(self):
        first_user = User.objects.create_user(username='first-client', first_name='First')
        second_user = User.objects.create_user(username='second-client', first_name='Second')
        self.first_client = ClientProfile.objects.create(user=first_user, id_number='FIT-001')
        self.second_client = ClientProfile.objects.create(user=second_user, id_number='FIT-002')
        self.training_class = TrainingClass.objects.create(
            name='Morning Boxing',
            class_type=FitnessClassType.BOXING,
        )

    def test_default_price_is_100_mad(self):
        self.assertEqual(self.training_class.price_per_member, Decimal('100.00'))
        self.assertEqual(self.training_class.team_total, Decimal('0.00'))

    def test_team_total_is_automatic(self):
        ClassMember.objects.create(training_class=self.training_class, client=self.first_client)
        self.assertEqual(self.training_class.team_total, Decimal('100.00'))

        ClassMember.objects.create(training_class=self.training_class, client=self.second_client)
        self.assertEqual(self.training_class.team_total, Decimal('200.00'))

        ClassMember.objects.filter(client=self.first_client).delete()
        self.assertEqual(self.training_class.team_total, Decimal('100.00'))


class GymClassCrudApiTest(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(username='class-admin', is_staff=True)
        self.client.force_login(self.admin)

    def test_reception_can_list_but_cannot_change_classes(self):
        TrainingClass.objects.create(name='Morning Boxing', class_type=FitnessClassType.BOXING)
        reception = User.objects.create_user(username='class-reception')
        Group.objects.get_or_create(name='Reception')[0].user_set.add(reception)
        self.client.force_login(reception)
        listed = self.client.get('/api/fitness/classes')
        self.assertEqual(listed.status_code, 200)
        self.assertEqual(len(listed.json()), 1)
        payload = json.dumps({
            'name': 'Evening Kick',
            'class_type': 'kick_boxing',
            'price_per_member': '120.00',
        })
        self.assertEqual(
            self.client.post('/api/fitness/classes', data=payload, content_type='application/json').status_code,
            403,
        )
        class_id = listed.json()[0]['id']
        self.assertEqual(
            self.client.put(
                f'/api/fitness/classes/{class_id}',
                data=payload,
                content_type='application/json',
            ).status_code,
            403,
        )
        self.assertEqual(self.client.delete(f'/api/fitness/classes/{class_id}').status_code, 403)

    def test_admin_can_add_edit_and_remove_a_class(self):
        created = self.client.post(
            '/api/fitness/classes',
            data=json.dumps({
                'name': 'Morning Boxing',
                'class_type': 'boxing',
                'price_per_member': '150.00',
                'is_active': True,
            }),
            content_type='application/json',
        )
        self.assertEqual(created.status_code, 200)
        class_id = created.json()['id']
        self.assertEqual(created.json()['name'], 'Morning Boxing')
        self.assertEqual(Decimal(str(created.json()['price_per_member'])), Decimal('150.00'))

        updated = self.client.put(
            f'/api/fitness/classes/{class_id}',
            data=json.dumps({
                'name': 'Evening Boxing',
                'class_type': 'boxing',
                'price_per_member': '180.00',
                'is_active': False,
            }),
            content_type='application/json',
        )
        self.assertEqual(updated.status_code, 200)
        self.assertEqual(updated.json()['name'], 'Evening Boxing')
        self.assertEqual(Decimal(str(updated.json()['price_per_member'])), Decimal('180.00'))
        self.assertFalse(updated.json()['is_active'])

        member_user = User.objects.create_user(username='class-assigned', first_name='Sara')
        member = ClientProfile.objects.create(user=member_user, id_number='CLS-001')
        ClassMember.objects.create(training_class_id=class_id, client=member)
        deleted = self.client.delete(f'/api/fitness/classes/{class_id}')
        self.assertEqual(deleted.status_code, 200)
        self.assertFalse(TrainingClass.objects.filter(id=class_id).exists())
        self.assertFalse(ClassMember.objects.filter(client=member).exists())
        self.assertTrue(ClientProfile.objects.filter(id=member.id).exists())

    def test_duplicate_class_name_and_type_is_rejected(self):
        TrainingClass.objects.create(name='Morning Boxing', class_type=FitnessClassType.BOXING)
        duplicate = self.client.post(
            '/api/fitness/classes',
            data=json.dumps({
                'name': 'Morning Boxing',
                'class_type': 'boxing',
                'price_per_member': '100.00',
            }),
            content_type='application/json',
        )
        self.assertEqual(duplicate.status_code, 409)


class GymPlanCrudApiTest(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(username='plan-admin', is_staff=True)
        self.client.force_login(self.admin)

    def test_reception_can_list_but_cannot_change_plans(self):
        MembershipPlan.objects.create(name='Monthly', duration_months=1, price=Decimal('300.00'))
        reception = User.objects.create_user(username='plan-reception')
        Group.objects.get_or_create(name='Reception')[0].user_set.add(reception)
        self.client.force_login(reception)
        listed = self.client.get('/api/fitness/plans')
        self.assertEqual(listed.status_code, 200)
        self.assertEqual(len(listed.json()), 1)
        payload = json.dumps({
            'name': 'Quarterly',
            'duration_months': 3,
            'price': '800.00',
            'description': '',
            'is_active': True,
        })
        self.assertEqual(
            self.client.post('/api/fitness/plans', data=payload, content_type='application/json').status_code,
            403,
        )
        plan_id = listed.json()[0]['id']
        self.assertEqual(
            self.client.put(
                f'/api/fitness/plans/{plan_id}',
                data=payload,
                content_type='application/json',
            ).status_code,
            403,
        )
        self.assertEqual(self.client.delete(f'/api/fitness/plans/{plan_id}').status_code, 403)

    def test_admin_can_add_edit_deactivate_and_remove_a_plan(self):
        created = self.client.post(
            '/api/fitness/plans',
            data=json.dumps({
                'name': 'Monthly',
                'duration_months': 1,
                'price': '300.00',
                'description': 'One month access',
                'is_active': True,
            }),
            content_type='application/json',
        )
        self.assertEqual(created.status_code, 200)
        plan_id = created.json()['id']
        self.assertEqual(created.json()['name'], 'Monthly')
        self.assertEqual(Decimal(str(created.json()['price'])), Decimal('300.00'))
        self.assertEqual(created.json()['member_count'], 0)

        updated = self.client.put(
            f'/api/fitness/plans/{plan_id}',
            data=json.dumps({
                'name': 'Monthly Plus',
                'duration_months': 1,
                'price': '350.00',
                'description': 'Updated',
                'is_active': False,
            }),
            content_type='application/json',
        )
        self.assertEqual(updated.status_code, 200)
        self.assertEqual(updated.json()['name'], 'Monthly Plus')
        self.assertEqual(Decimal(str(updated.json()['price'])), Decimal('350.00'))
        self.assertFalse(updated.json()['is_active'])

        listed = self.client.get('/api/fitness/plans')
        self.assertEqual(len(listed.json()), 1)
        self.assertFalse(listed.json()[0]['is_active'])

        deleted = self.client.delete(f'/api/fitness/plans/{plan_id}')
        self.assertEqual(deleted.status_code, 200)
        self.assertFalse(MembershipPlan.objects.filter(id=plan_id).exists())

    def test_cannot_delete_a_plan_that_has_memberships(self):
        plan = MembershipPlan.objects.create(name='Monthly', duration_months=1, price=Decimal('300.00'))
        member_user = User.objects.create_user(username='plan-member', first_name='Sara')
        member = ClientProfile.objects.create(user=member_user, id_number='PLAN-001')
        Membership.objects.create(
            member=member,
            plan=plan,
            start_date='2026-01-01',
            end_date='2026-02-01',
            price=plan.price,
        )
        listed = self.client.get('/api/fitness/plans')
        self.assertEqual(listed.json()[0]['member_count'], 1)
        blocked = self.client.delete(f'/api/fitness/plans/{plan.id}')
        self.assertEqual(blocked.status_code, 400)
        self.assertTrue(MembershipPlan.objects.filter(id=plan.id).exists())

    def test_duplicate_plan_name_is_rejected(self):
        MembershipPlan.objects.create(name='Monthly', duration_months=1, price=Decimal('300.00'))
        duplicate = self.client.post(
            '/api/fitness/plans',
            data=json.dumps({
                'name': 'Monthly',
                'duration_months': 3,
                'price': '800.00',
                'description': '',
                'is_active': True,
            }),
            content_type='application/json',
        )
        self.assertEqual(duplicate.status_code, 409)


class MembershipCrudApiTest(TestCase):
    def setUp(self):
        self.staff = User.objects.create_user(username='member-staff', is_staff=True)
        self.client.force_login(self.staff)
        self.user = User.objects.create_user(username='member-user', first_name='Member', last_name='User')
        self.client_profile = ClientProfile.objects.create(user=self.user, phone='0612345678', id_number='FIT-010')
        self.plan_one = self._create_plan('Starter', 1, Decimal('300.00'))
        self.plan_two = self._create_plan('Premium', 3, Decimal('700.00'))
        self.membership = self._create_membership(self.plan_one, self.client_profile, '2026-01-01')

    def _create_plan(self, name, duration, price):
        from fitness.models import MembershipPlan
        return MembershipPlan.objects.create(name=name, duration_months=duration, price=price)

    def _create_membership(self, plan, member, start_date):
        from fitness.models import Membership
        return Membership.objects.create(
            member=member,
            plan=plan,
            start_date=start_date,
            end_date='2026-02-01',
            price=plan.price,
            notes='Initial plan',
        )

    def test_update_membership_endpoint(self):
        payload = {
            'member_id': self.client_profile.id,
            'plan_id': self.plan_two.id,
            'start_date': '2026-02-10',
            'notes': 'Updated plan',
        }

        response = self.client.put(
            f'/api/fitness/memberships/{self.membership.id}',
            data=json.dumps(payload),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 200)
        self.membership.refresh_from_db()
        self.assertEqual(self.membership.plan_id, self.plan_two.id)
        self.assertEqual(self.membership.notes, 'Updated plan')
        self.assertEqual(str(self.membership.price), '700.00')

    def test_update_membership_can_set_a_custom_price(self):
        payload = {
            'member_id': self.client_profile.id,
            'plan_id': self.plan_one.id,
            'start_date': '2026-01-01',
            'notes': 'Custom price',
            'price': '250.00',
        }

        response = self.client.put(
            f'/api/fitness/memberships/{self.membership.id}',
            data=json.dumps(payload),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 200)
        self.membership.refresh_from_db()
        self.assertEqual(str(self.membership.price), '250.00')

    def test_membership_price_can_be_patched(self):
        response = self.client.patch(
            f'/api/fitness/memberships/{self.membership.id}/price',
            data=json.dumps({'price': '180.00'}),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)
        self.membership.refresh_from_db()
        self.assertEqual(str(self.membership.price), '180.00')
        self.assertEqual(Decimal(str(response.json()['price'])), Decimal('180.00'))

    def test_delete_membership_endpoint(self):
        response = self.client.delete(f'/api/fitness/memberships/{self.membership.id}')

        self.assertEqual(response.status_code, 200)
        self.assertFalse(self.membership.__class__.objects.filter(id=self.membership.id).exists())

    def test_record_payment_returns_400_when_amount_exceeds_balance(self):
        response = self.client.post(
            f'/api/fitness/memberships/{self.membership.id}/payments',
            data=json.dumps({'amount': '999.00', 'received_by': 'Admin', 'notes': ''}),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()['detail'], 'Payment exceeds remaining balance')

    def test_record_payment_endpoint(self):
        list_response = self.client.get(f'/api/fitness/memberships/{self.membership.id}/payments')
        self.assertEqual(list_response.status_code, 200)

        response = self.client.post(
            f'/api/fitness/memberships/{self.membership.id}/payments',
            data=json.dumps({'amount': '50.00', 'received_by': 'Admin', 'notes': 'Cash'}),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Decimal(str(response.json()['amount'])), Decimal('50.00'))

    def test_record_payment_can_set_remaining_balance(self):
        response = self.client.post(
            f'/api/fitness/memberships/{self.membership.id}/payments',
            data=json.dumps({'amount': '100.00', 'received_by': 'Admin', 'notes': 'Cash', 'remaining': '20.00'}),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)
        self.membership.refresh_from_db()
        self.assertEqual(self.membership.price, Decimal('120.00'))
        self.assertEqual(self.membership.total_paid, Decimal('100.00'))
        self.assertEqual(self.membership.remaining_balance, Decimal('20.00'))

    def test_membership_remaining_can_be_patched(self):
        GymPayment.objects.create(membership=self.membership, amount=Decimal('80.00'), received_by='Admin')
        response = self.client.patch(
            f'/api/fitness/memberships/{self.membership.id}/remaining',
            data=json.dumps({'remaining': '20.00'}),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Decimal(str(response.json()['price'])), Decimal('100.00'))
        self.assertEqual(Decimal(str(response.json()['remaining_balance'])), Decimal('20.00'))


class MemberClassApiTest(TestCase):
    def setUp(self):
        self.staff = User.objects.create_user(username='class-staff', is_staff=True)
        self.client.force_login(self.staff)
        self.user = User.objects.create_user(username='class-member-user', first_name='Class', last_name='Member')
        self.member = ClientProfile.objects.create(user=self.user, phone='0699999999', id_number='FIT-CLASS')
        self.boxing = TrainingClass.objects.create(name='Boxing Team', class_type=FitnessClassType.BOXING)
        self.aerobic = TrainingClass.objects.create(name='Aerobic Team', class_type=FitnessClassType.AEROBIC)

    def test_assign_and_switch_member_class(self):
        empty = self.client.get(f'/api/fitness/members/{self.member.id}/class')
        self.assertEqual(empty.status_code, 200)
        self.assertIsNone(empty.json()['training_class_id'])

        assigned = self.client.put(
            f'/api/fitness/members/{self.member.id}/class',
            data=json.dumps({'class_id': self.boxing.id}),
            content_type='application/json',
        )
        self.assertEqual(assigned.status_code, 200)
        self.assertEqual(assigned.json()['training_class_id'], self.boxing.id)
        self.assertTrue(ClassMember.objects.get(training_class=self.boxing, client=self.member).is_active)

        switched = self.client.put(
            f'/api/fitness/members/{self.member.id}/class',
            data=json.dumps({'class_id': self.aerobic.id}),
            content_type='application/json',
        )
        self.assertEqual(switched.status_code, 200)
        self.assertEqual(switched.json()['training_class_id'], self.aerobic.id)
        self.assertFalse(ClassMember.objects.get(training_class=self.boxing, client=self.member).is_active)
        self.assertTrue(ClassMember.objects.get(training_class=self.aerobic, client=self.member).is_active)

        cleared = self.client.put(
            f'/api/fitness/members/{self.member.id}/class',
            data=json.dumps({'class_id': None}),
            content_type='application/json',
        )
        self.assertEqual(cleared.status_code, 200)
        self.assertIsNone(cleared.json()['training_class_id'])
        self.assertFalse(ClassMember.objects.filter(client=self.member, is_active=True).exists())


class MemberArchiveApiTest(TestCase):
    def test_archiving_member_removes_active_class_assignment(self):
        staff = User.objects.create_user(username='archive-staff', is_staff=True)
        self.client.force_login(staff)
        user = User.objects.create_user(username='archive-member', first_name='Archive')
        member = ClientProfile.objects.create(user=user, id_number='FIT-ARCHIVE')
        training_class = TrainingClass.objects.create(name='Archive Boxing', class_type=FitnessClassType.BOXING)
        ClassMember.objects.create(training_class=training_class, client=member)
        self.assertEqual(training_class.member_count, 1)

        response = self.client.delete(f'/api/fitness/members/{member.id}')

        self.assertEqual(response.status_code, 200)
        member.refresh_from_db()
        self.assertFalse(member.is_active)
        self.assertFalse(ClassMember.objects.get(training_class=training_class, client=member).is_active)
        self.assertEqual(training_class.member_count, 0)

    def test_archived_member_is_not_counted_in_class(self):
        user = User.objects.create_user(username='inactive-class-member')
        member = ClientProfile.objects.create(user=user, id_number='FIT-INACTIVE')
        training_class = TrainingClass.objects.create(name='Inactive Boxing', class_type=FitnessClassType.BOXING)
        ClassMember.objects.create(training_class=training_class, client=member)
        member.is_active = False
        member.save(update_fields=['is_active'])
        self.assertEqual(training_class.member_count, 0)


class MembershipExpiryNotificationTest(TestCase):
    def test_creates_one_notification_when_membership_expires_in_seven_days(self):
        admin = User.objects.create_user(username='expiry-admin', is_staff=True)
        member_user = User.objects.create_user(username='expiry-member', first_name='Expiry', last_name='Member')
        member = ClientProfile.objects.create(user=member_user, id_number='FIT-EXPIRY')
        plan = MembershipPlan.objects.create(name='Expiry Plan', duration_months=1, price=Decimal('120.00'))
        today = timezone.localdate()
        Membership.objects.create(member=member, plan=plan, start_date=today - timedelta(days=23), end_date=today + timedelta(days=7), price=plan.price)

        create_expiring_membership_notifications()
        create_expiring_membership_notifications()

        notifications = GymNotification.objects.filter(recipient=admin, member_id=member.id)
        self.assertEqual(notifications.count(), 1)
        self.assertIn('7 days remaining', notifications.first().message)


class GymNotificationApiTest(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(username='notify-admin', password='password123', is_staff=True, first_name='Notify', last_name='Admin')
        self.client.force_login(self.admin)

    def test_creating_a_member_creates_a_notification(self):
        response = self.client.post('/api/fitness/members', data=json.dumps({
            'first_name': 'Sara',
            'last_name': 'Benali',
            'phone': '0611111111',
            'email': 'sara@example.com',
            'id_number': 'BK123456',
            'address': '12 Rue des Fleurs',
            'city': 'Casablanca',
            'country': 'Morocco',
            'postal_code': '',
        }), content_type='application/json')
        self.assertEqual(response.status_code, 200)
        notifications = self.client.get('/api/notifications').json()
        self.assertTrue(any(item['title'] == 'New member registered' for item in notifications))

    def test_notification_settings_serialize_for_staff(self):
        response = self.client.get('/api/notifications/settings')
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertIn('id', payload)
        self.assertIn('membership_expiring_soon', payload)

    def test_notification_csrf_is_required_from_the_frontend_origin(self):
        from django.test import Client
        GymNotification.objects.create(recipient=self.admin, category='members', title='Test alert', message='Something happened.')
        notification_id = GymNotification.objects.get(recipient=self.admin).id
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.admin)
        denied = client.delete(
            f'/api/notifications/{notification_id}',
            HTTP_ORIGIN='http://localhost:5173',
        )
        self.assertEqual(denied.status_code, 403)
        self.assertTrue(GymNotification.objects.filter(id=notification_id).exists())

        client.get('/api/auth/me')
        token = client.cookies['csrftoken'].value
        allowed = client.delete(
            f'/api/notifications/{notification_id}',
            HTTP_ORIGIN='http://localhost:5173',
            HTTP_X_CSRFTOKEN=token,
        )
        self.assertEqual(allowed.status_code, 200)
        self.assertFalse(GymNotification.objects.filter(id=notification_id).exists())

    def test_all_notifications_can_be_deleted(self):
        GymNotification.objects.create(recipient=self.admin, category='members', title='First', message='One')
        GymNotification.objects.create(recipient=self.admin, category='payments', title='Second', message='Two')
        response = self.client.delete('/api/notifications')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(GymNotification.objects.filter(recipient=self.admin).count(), 0)

    def test_payment_status_change_creates_a_notification(self):
        member_user = User.objects.create_user(username='paid-member', first_name='Lina', last_name='Said')
        member = ClientProfile.objects.create(user=member_user, phone='0622222222', id_number='GYM-PAY')
        plan = MembershipPlan.objects.create(name='Monthly', duration_months=1, price=Decimal('400.00'))
        membership = Membership.objects.create(
            member=member,
            plan=plan,
            start_date=timezone.localdate(),
            end_date=timezone.localdate() + timedelta(days=30),
            price=plan.price,
        )
        response = self.client.patch(
            f'/api/fitness/memberships/{membership.id}/payment-status',
            data=json.dumps({'status': 'paid'}),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)
        notifications = self.client.get('/api/notifications').json()
        self.assertTrue(any(item['title'] == 'Payment marked as paid' for item in notifications))

    def test_reception_cannot_read_or_change_notification_settings(self):
        reception = User.objects.create_user(username='notify-reception')
        Group.objects.get_or_create(name='Reception')[0].user_set.add(reception)
        self.client.force_login(reception)
        self.assertEqual(self.client.get('/api/notifications/settings').status_code, 403)
        denied = self.client.put(
            '/api/notifications/settings',
            data=json.dumps({'membership_expiring_soon': False}),
            content_type='application/json',
        )
        self.assertEqual(denied.status_code, 403)

    def test_staff_admin_notifications_are_not_delivered_to_reception(self):
        reception = User.objects.create_user(username='staff-notify-reception')
        Group.objects.get_or_create(name='Reception')[0].user_set.add(reception)
        created = self.client.post(
            '/api/admin/users',
            data=json.dumps({
                'username': 'new-desk',
                'password': 'DeskPass-2026!',
                'first_name': 'New',
                'last_name': 'Desk',
                'email': 'newdesk@example.com',
                'role': 'Reception',
            }),
            content_type='application/json',
        )
        self.assertEqual(created.status_code, 200)
        self.assertTrue(GymNotification.objects.filter(recipient=self.admin, title='New staff user created').exists())
        self.assertFalse(GymNotification.objects.filter(recipient=reception, title='New staff user created').exists())

        GymNotification.objects.create(
            recipient=reception,
            category='system',
            title='User role changed',
            message='Should stay hidden',
        )
        self.client.force_login(reception)
        titles = [item['title'] for item in self.client.get('/api/notifications').json()]
        self.assertNotIn('New staff user created', titles)
        self.assertNotIn('User role changed', titles)


class ClassRevenueReportTest(TestCase):
    def test_class_report_calculates_monthly_revenue_per_class(self):
        admin = User.objects.create_user(username='report-admin', is_staff=True)
        self.client.force_login(admin)
        boxing = TrainingClass.objects.create(name='Morning Boxing', class_type=FitnessClassType.BOXING, price_per_member=Decimal('200.00'))
        aerobic = TrainingClass.objects.create(name='Aerobic', class_type=FitnessClassType.AEROBIC, price_per_member=Decimal('150.00'))
        plan = MembershipPlan.objects.create(name='Report Plan', duration_months=1, price=Decimal('200.00'))
        today = timezone.localdate()

        first_user = User.objects.create_user(username='boxer-one', first_name='Amine')
        first_member = ClientProfile.objects.create(user=first_user, id_number='REP-1')
        ClassMember.objects.create(training_class=boxing, client=first_member)
        first_membership = Membership.objects.create(
            member=first_member,
            plan=plan,
            start_date=today.replace(day=1),
            end_date=today.replace(day=1) + timedelta(days=30),
            price=Decimal('200.00'),
        )
        GymPayment.objects.create(membership=first_membership, amount=Decimal('200.00'), received_by='Admin')

        second_user = User.objects.create_user(username='aero-one', first_name='Sara')
        second_member = ClientProfile.objects.create(user=second_user, id_number='REP-2')
        ClassMember.objects.create(training_class=aerobic, client=second_member)
        second_membership = Membership.objects.create(
            member=second_member,
            plan=plan,
            start_date=today.replace(day=1),
            end_date=today.replace(day=1) + timedelta(days=30),
            price=Decimal('150.00'),
        )
        GymPayment.objects.create(membership=second_membership, amount=Decimal('50.00'), received_by='Admin')

        response = self.client.get(f'/api/fitness/reports/classes?year={today.year}&month={today.month}')
        self.assertEqual(response.status_code, 200)
        body = response.json()
        by_name = {item['name']: item for item in body['classes']}
        self.assertEqual(by_name['Morning Boxing']['member_count'], 1)
        self.assertEqual(Decimal(str(by_name['Morning Boxing']['expected_monthly'])), Decimal('200.00'))
        self.assertEqual(Decimal(str(by_name['Morning Boxing']['collected'])), Decimal('200.00'))
        self.assertEqual(Decimal(str(by_name['Aerobic']['expected_monthly'])), Decimal('150.00'))
        self.assertEqual(Decimal(str(by_name['Aerobic']['collected'])), Decimal('50.00'))
        self.assertEqual(Decimal(str(by_name['Aerobic']['outstanding'])), Decimal('100.00'))
        self.assertEqual(Decimal(str(body['total_collected'])), Decimal('250.00'))


class TrainerPayrollApiTest(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(username='trainer-admin', is_staff=True)
        self.client.force_login(self.admin)

    def test_only_admin_can_add_a_trainer(self):
        reception = User.objects.create_user(username='trainer-reception')
        Group.objects.get_or_create(name='Reception')[0].user_set.add(reception)
        self.client.force_login(reception)
        denied = self.client.post(
            '/api/fitness/trainers',
            data=json.dumps({'first_name': 'Karim', 'last_name': 'Ali'}),
            content_type='application/json',
        )
        self.assertEqual(denied.status_code, 403)
        self.assertEqual(self.client.get('/api/fitness/trainers').status_code, 403)

        self.client.force_login(self.admin)
        created = self.client.post(
            '/api/fitness/trainers',
            data=json.dumps({
                'first_name': 'Karim',
                'last_name': 'Ali',
                'specialization': 'Boxing',
                'monthly_pay': '400.00',
            }),
            content_type='application/json',
        )
        self.assertEqual(created.status_code, 200)
        body = created.json()
        self.assertEqual(body['first_name'], 'Karim')
        self.assertEqual(Decimal(str(body['monthly_pay'])), Decimal('400.00'))
        self.assertEqual(Decimal(str(body['pay_amount'])), Decimal('400.00'))
        self.assertFalse(body['is_paid'])

    def test_payroll_can_be_marked_paid_and_appears_in_reports(self):
        trainer = Trainer.objects.create(first_name='Nadia', last_name='Ben', monthly_pay=Decimal('400.00'))
        today = timezone.localdate()
        paid = self.client.patch(
            f'/api/fitness/trainers/{trainer.id}/payroll',
            data=json.dumps({
                'year': today.year,
                'month': today.month,
                'pay_amount': '400.00',
                'is_paid': True,
            }),
            content_type='application/json',
        )
        self.assertEqual(paid.status_code, 200)
        self.assertTrue(paid.json()['is_paid'])
        self.assertTrue(TrainerPayroll.objects.get(trainer=trainer, year=today.year, month=today.month).is_paid)

        report = self.client.get(f'/api/fitness/reports/trainers?year={today.year}&month={today.month}')
        self.assertEqual(report.status_code, 200)
        body = report.json()
        self.assertEqual(Decimal(str(body['total_paid'])), Decimal('400.00'))
        self.assertEqual(Decimal(str(body['total_unpaid'])), Decimal('0.00'))
        self.assertEqual(body['trainers'][0]['name'], 'Nadia Ben')
        self.assertTrue(body['trainers'][0]['is_paid'])

    def test_admin_can_delete_a_trainer(self):
        trainer = Trainer.objects.create(first_name='Omar', last_name='Said', monthly_pay=Decimal('350.00'))
        removed = self.client.delete(f'/api/fitness/trainers/{trainer.id}')
        self.assertEqual(removed.status_code, 200)
        trainer.refresh_from_db()
        self.assertFalse(trainer.is_active)
        listed = self.client.get('/api/fitness/trainers').json()
        self.assertFalse(any(item['id'] == trainer.id for item in listed))

    def test_trainer_notifications_go_only_to_admins(self):
        reception = User.objects.create_user(username='trainer-reception-notify', is_staff=True)
        Group.objects.get_or_create(name='Reception')[0].user_set.add(reception)
        created = self.client.post(
            '/api/fitness/trainers',
            data=json.dumps({'first_name': 'Karim', 'last_name': 'Ali', 'monthly_pay': '400.00'}),
            content_type='application/json',
        )
        self.assertEqual(created.status_code, 200)
        self.assertTrue(GymNotification.objects.filter(recipient=self.admin, title='Trainer added').exists())
        self.assertFalse(GymNotification.objects.filter(recipient=reception, title='Trainer added').exists())

        GymNotification.objects.create(
            recipient=reception,
            category='system',
            title='Trainer payroll updated',
            message='Old trainer alert',
        )
        self.client.force_login(reception)
        titles = [item['title'] for item in self.client.get('/api/notifications').json()]
        self.assertNotIn('Trainer added', titles)
        self.assertNotIn('Trainer payroll updated', titles)


class GymExpenseApiTest(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(username='expense-admin', is_staff=True)
        self.client.force_login(self.admin)
        self.today = timezone.localdate()

    def test_only_admin_can_manage_expenses(self):
        reception = User.objects.create_user(username='expense-reception')
        Group.objects.get_or_create(name='Reception')[0].user_set.add(reception)
        self.client.force_login(reception)
        denied = self.client.post(
            '/api/fitness/expenses',
            data=json.dumps({'category': 'water', 'title': 'Water bill', 'amount': '80.00'}),
            content_type='application/json',
        )
        self.assertEqual(denied.status_code, 403)
        self.assertEqual(self.client.get('/api/fitness/expenses').status_code, 403)

    def test_admin_can_record_monthly_expenses_and_see_overview(self):
        Trainer.objects.create(first_name='Karim', last_name='Ali', monthly_pay=Decimal('400.00'))
        created = self.client.post(
            '/api/fitness/expenses',
            data=json.dumps({
                'category': 'electricity',
                'title': 'ONEE bill',
                'amount': '250.00',
                'year': self.today.year,
                'month': self.today.month,
            }),
            content_type='application/json',
        )
        self.assertEqual(created.status_code, 200)
        self.assertEqual(created.json()['category'], 'electricity')
        self.assertEqual(Decimal(str(created.json()['amount'])), Decimal('250.00'))

        water = self.client.post(
            '/api/fitness/expenses',
            data=json.dumps({
                'category': 'cleaning',
                'title': 'Cleaning lady',
                'amount': '150.00',
                'year': self.today.year,
                'month': self.today.month,
            }),
            content_type='application/json',
        )
        self.assertEqual(water.status_code, 200)

        listed = self.client.get(
            f'/api/fitness/expenses?year={self.today.year}&month={self.today.month}'
        )
        self.assertEqual(listed.status_code, 200)
        self.assertEqual(len(listed.json()), 2)

        overview = self.client.get(
            f'/api/fitness/reports/overview?year={self.today.year}&month={self.today.month}'
        )
        self.assertEqual(overview.status_code, 200)
        body = overview.json()
        self.assertEqual(Decimal(str(body['operating_total'])), Decimal('400.00'))
        self.assertEqual(Decimal(str(body['trainer_due'])), Decimal('400.00'))
        self.assertEqual(Decimal(str(body['total_spend'])), Decimal('800.00'))
        self.assertEqual(Decimal(str(body['collected'])), Decimal('0.00'))
        self.assertEqual(Decimal(str(body['net'])), Decimal('-800.00'))
        self.assertEqual(len(body['categories']), 2)

        expense_id = created.json()['id']
        deleted = self.client.delete(f'/api/fitness/expenses/{expense_id}')
        self.assertEqual(deleted.status_code, 200)
        self.assertEqual(GymExpense.objects.count(), 1)

    def test_expense_notifications_go_only_to_admins(self):
        reception = User.objects.create_user(username='expense-reception-notify', is_staff=True)
        Group.objects.get_or_create(name='Reception')[0].user_set.add(reception)
        created = self.client.post(
            '/api/fitness/expenses',
            data=json.dumps({
                'category': 'water',
                'title': 'Water bill',
                'amount': '80.00',
                'year': self.today.year,
                'month': self.today.month,
            }),
            content_type='application/json',
        )
        self.assertEqual(created.status_code, 200)
        self.assertTrue(GymNotification.objects.filter(recipient=self.admin, title='Expense recorded').exists())
        self.assertFalse(GymNotification.objects.filter(recipient=reception, title='Expense recorded').exists())

        GymNotification.objects.create(
            recipient=reception,
            category='system',
            title='Expense recorded',
            message='Old expense alert',
        )
        self.client.force_login(reception)
        titles = [item['title'] for item in self.client.get('/api/notifications').json()]
        self.assertNotIn('Expense recorded', titles)


class MonthlyReportExportTest(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(username='export-admin', is_staff=True)
        self.client.force_login(self.admin)
        self.today = timezone.localdate()
        Trainer.objects.create(first_name='Karim', last_name='Ali', monthly_pay=Decimal('400.00'))
        GymExpense.objects.create(
            category='electricity',
            title='ONEE bill',
            amount=Decimal('250.00'),
            year=self.today.year,
            month=self.today.month,
        )
        member_user = User.objects.create_user(username='export-member', first_name='Sara', last_name='Benali')
        member = ClientProfile.objects.create(user=member_user, phone='0612345678', id_number='EX-001')
        plan = MembershipPlan.objects.create(name='Monthly', duration_months=1, price=Decimal('400.00'))
        membership = Membership.objects.create(
            member=member,
            plan=plan,
            start_date=self.today.replace(day=1),
            end_date=self.today,
            price=Decimal('400.00'),
        )
        GymPayment.objects.create(membership=membership, amount=Decimal('100.00'), received_by='Admin')

    def _export_url(self, kind):
        return f'/api/fitness/reports/export/{kind}?year={self.today.year}&month={self.today.month}'

    def test_reception_cannot_download_monthly_report(self):
        reception = User.objects.create_user(username='export-reception')
        Group.objects.get_or_create(name='Reception')[0].user_set.add(reception)
        self.client.force_login(reception)
        self.assertEqual(self.client.get(self._export_url('xlsx')).status_code, 403)
        self.assertEqual(self.client.get(self._export_url('pdf')).status_code, 403)

    def test_admin_can_download_excel_with_income_bills_trainer_pay_and_net(self):
        from io import BytesIO
        from openpyxl import load_workbook

        response = self.client.get(self._export_url('xlsx'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertIn(
            f'AUMB-monthly-report-{self.today.year}-{self.today.month:02d}.xlsx',
            response['Content-Disposition'],
        )
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook['Summary']['A1'].value, 'AUMB')
        self.assertEqual(workbook.sheetnames, ['Summary', 'Income', 'Bills', 'Trainer pay'])
        amounts = {
            workbook['Summary'].cell(row, 1).value: workbook['Summary'].cell(row, 2).value
            for row in range(8, 16)
        }
        self.assertEqual(amounts['Cash collected'], 100.0)
        self.assertEqual(amounts['Operating expenses (bills)'], 250.0)
        self.assertEqual(amounts['Trainer pay due'], 400.0)
        self.assertEqual(amounts['Total spend'], 650.0)
        self.assertEqual(amounts['Net'], -550.0)
        self.assertEqual(workbook['Bills']['A2'].value, 'Electricity')
        self.assertEqual(workbook['Bills']['C2'].value, 250.0)
        self.assertEqual(workbook['Trainer pay']['A2'].value, 'Karim Ali')
        self.assertEqual(workbook['Trainer pay']['D2'].value, 400.0)

    def test_admin_can_download_pdf_monthly_report(self):
        response = self.client.get(self._export_url('pdf'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertTrue(response.content.startswith(b'%PDF'))
        self.assertIn(
            f'AUMB-monthly-report-{self.today.year}-{self.today.month:02d}.pdf',
            response['Content-Disposition'],
        )


class WhatsAppReminderApiTest(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(username='wa-admin', is_staff=True)
        self.client.force_login(self.admin)
        member_user = User.objects.create_user(username='wa-member', first_name='Sara', last_name='Benali')
        self.member = ClientProfile.objects.create(user=member_user, phone='0612345678', id_number='WA-001')
        self.plan = MembershipPlan.objects.create(name='Monthly', duration_months=1, price=Decimal('400.00'))
        today = timezone.localdate()
        self.membership = Membership.objects.create(
            member=self.member,
            plan=self.plan,
            start_date=today - timedelta(days=25),
            end_date=today + timedelta(days=5),
            price=Decimal('400.00'),
        )
        GymPayment.objects.create(membership=self.membership, amount=Decimal('100.00'), received_by='Admin')

    def test_lists_expiring_and_unpaid_members_with_whatsapp_link(self):
        response = self.client.get('/api/fitness/reminders')
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body['expiring'], 1)
        self.assertEqual(body['unpaid'], 1)
        row = body['items'][0]
        self.assertEqual(row['member_name'], 'Sara Benali')
        self.assertIn('expiring_soon', row['reasons'])
        self.assertIn('unpaid', row['reasons'])
        self.assertTrue(row['whatsapp_url'].startswith('https://wa.me/212612345678?text='))
        self.assertIn('expire', row['message'])
        self.assertIn('300.00 MAD', row['message'])
        dashboard = self.client.get('/api/fitness/dashboard').json()
        self.assertEqual(dashboard['whatsapp_due'], 1)

    def test_marking_a_reminder_sent_flags_it_for_today(self):
        sent = self.client.post('/api/fitness/reminders/%s/sent' % self.membership.id, data=json.dumps({}), content_type='application/json')
        self.assertEqual(sent.status_code, 200)
        self.assertTrue(sent.json()['reminded_today'])
        self.assertEqual(GymWhatsAppReminder.objects.filter(membership=self.membership).count(), 1)

    def test_ends_today_uses_month_end_reminder_message(self):
        Membership.objects.filter(id=self.membership.id).update(end_date=timezone.localdate())
        body = self.client.get('/api/fitness/reminders').json()
        self.assertEqual(len(body['items']), 1)
        message = body['items'][0]['message']
        self.assertIn('Bonjour Sara Benali', message)
        self.assertIn('se termine aujourd', message)
        self.assertIn('renouveler', message)
        self.assertIn('AUMB', message)
        self.assertIn('300.00 MAD', message)

    def test_skips_paid_active_memberships(self):
        Membership.objects.filter(id=self.membership.id).update(end_date=timezone.localdate() + timedelta(days=20))
        GymPayment.objects.create(membership=self.membership, amount=Decimal('300.00'), received_by='Admin')
        body = self.client.get('/api/fitness/reminders').json()
        self.assertEqual(body['items'], [])


class MemberIdentityApiTest(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(username='cin-admin', is_staff=True)
        self.client.force_login(self.admin)

    def test_create_member_requires_cin_and_address(self):
        missing = self.client.post(
            '/api/fitness/members',
            data=json.dumps({
                'first_name': 'Sara',
                'last_name': 'Benali',
                'phone': '0611111111',
                'email': 'sara@example.com',
            }),
            content_type='application/json',
        )
        self.assertEqual(missing.status_code, 400)

        created = self.client.post(
            '/api/fitness/members',
            data=json.dumps({
                'first_name': 'Sara',
                'last_name': 'Benali',
                'phone': '0611111111',
                'email': 'sara@example.com',
                'id_number': 'bk 123456',
                'address': '12 Rue des Fleurs',
                'city': 'Casablanca',
            }),
            content_type='application/json',
        )
        self.assertEqual(created.status_code, 200)
        self.assertEqual(created.json()['id_number'], 'BK123456')
        self.assertEqual(created.json()['address'], '12 Rue des Fleurs')
        self.assertEqual(created.json()['city'], 'Casablanca')

        duplicate = self.client.post(
            '/api/fitness/members',
            data=json.dumps({
                'first_name': 'Lina',
                'last_name': 'Said',
                'phone': '0622222222',
                'email': 'lina@example.com',
                'id_number': 'BK123456',
                'address': 'Another street',
                'city': 'Rabat',
            }),
            content_type='application/json',
        )
        self.assertEqual(duplicate.status_code, 409)

    def test_members_can_be_found_by_cin(self):
        self.client.post(
            '/api/fitness/members',
            data=json.dumps({
                'first_name': 'Sara',
                'last_name': 'Benali',
                'id_number': 'AB987654',
                'address': 'Hay Mohammadi',
                'city': 'Casablanca',
            }),
            content_type='application/json',
        )
        found = self.client.get('/api/fitness/members?search=AB987654')
        self.assertEqual(found.status_code, 200)
        self.assertEqual(found.json()[0]['id_number'], 'AB987654')


class AttendanceDeskApiTest(TestCase):
    def setUp(self):
        self.staff = User.objects.create_user(username='desk-staff', is_staff=True)
        self.client.force_login(self.staff)
        today = timezone.localdate()
        member_user = User.objects.create_user(username='desk-member', first_name='Sara', last_name='Benali')
        self.member = ClientProfile.objects.create(user=member_user, phone='0612345678', id_number='CARD-9')
        self.boxing = TrainingClass.objects.create(name='Boxing Team', class_type=FitnessClassType.BOXING)
        ClassMember.objects.create(training_class=self.boxing, client=self.member)
        plan = MembershipPlan.objects.create(name='Desk plan', duration_months=1, price=Decimal('400.00'))
        Membership.objects.create(
            member=self.member,
            plan=plan,
            start_date=today - timedelta(days=2),
            end_date=today + timedelta(days=20),
            price=Decimal('400.00'),
        )

    def test_lookup_by_name_phone_card_and_id_number(self):
        card = f'FO-{self.member.id:06d}'
        for query in ('Sara', '0612345678', card, 'CARD-9', f'flexoper:member:{self.member.id}'):
            response = self.client.get(f'/api/fitness/attendance/lookup?q={query}')
            self.assertEqual(response.status_code, 200, query)
            body = response.json()
            self.assertEqual(body['matches'][0]['id'], self.member.id, query)
            self.assertEqual(body['matches'][0]['card_code'], card)

    def test_check_in_check_out_and_class_headcount(self):
        created = self.client.post(
            '/api/fitness/attendance/check-in',
            data=json.dumps({'member_id': self.member.id}),
            content_type='application/json',
        )
        self.assertEqual(created.status_code, 200)
        self.assertEqual(created.json()['member_name'], 'Sara Benali')
        self.assertEqual(created.json()['class_name'], 'Boxing Team')
        self.assertTrue(created.json()['is_inside'])

        duplicate = self.client.post(
            '/api/fitness/attendance/check-in',
            data=json.dumps({'member_id': self.member.id}),
            content_type='application/json',
        )
        self.assertEqual(duplicate.status_code, 409)

        desk = self.client.get('/api/fitness/attendance/desk').json()
        self.assertEqual(desk['inside'], 1)
        self.assertEqual(desk['checkins'], 1)
        boxing = next(item for item in desk['by_class'] if item['class_id'] == self.boxing.id)
        self.assertEqual(boxing['inside'], 1)
        self.assertEqual(boxing['checkins'], 1)

        left = self.client.post(
            '/api/fitness/attendance/check-out',
            data=json.dumps({'member_id': self.member.id}),
            content_type='application/json',
        )
        self.assertEqual(left.status_code, 200)
        self.assertFalse(left.json()['is_inside'])
        self.assertIsNotNone(left.json()['checked_out_at'])

        again = self.client.post(
            '/api/fitness/attendance/check-in',
            data=json.dumps({'member_id': self.member.id}),
            content_type='application/json',
        )
        self.assertEqual(again.status_code, 200)
        self.assertNotEqual(again.json()['id'], created.json()['id'])

    def test_membership_card_qr_is_available(self):
        response = self.client.get(f'/api/fitness/members/{self.member.id}/qr')
        self.assertEqual(response.status_code, 200)
        self.assertIn('image/svg', response['Content-Type'])
        self.assertIn(b'<svg', response.content)

    def test_member_without_active_membership_cannot_check_in(self):
        other_user = User.objects.create_user(username='desk-guest', first_name='Guest')
        guest = ClientProfile.objects.create(user=other_user, id_number='GUEST-1')
        denied = self.client.post(
            '/api/fitness/attendance/check-in',
            data=json.dumps({'member_id': guest.id}),
            content_type='application/json',
        )
        self.assertEqual(denied.status_code, 400)


class GymCashDeskApiTest(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(username='cash-admin', is_staff=True)
        self.client.force_login(self.admin)
        self.today = timezone.localdate()
        member_user = User.objects.create_user(username='cash-member', first_name='Sara', last_name='Benali')
        self.member = ClientProfile.objects.create(user=member_user, phone='0612345678', id_number='CIN-CASH')
        plan = MembershipPlan.objects.create(name='Monthly', duration_months=1, price=Decimal('400.00'))
        self.membership = Membership.objects.create(
            member=self.member,
            plan=plan,
            start_date=self.today.replace(day=1),
            end_date=self.today,
            price=Decimal('400.00'),
        )
        self.payment = GymPayment.objects.create(
            membership=self.membership,
            amount=Decimal('150.00'),
            received_by='Admin',
            notes='Front desk cash',
        )

    def test_payment_list_includes_member_and_receipt(self):
        response = self.client.get('/api/fitness/payments')
        self.assertEqual(response.status_code, 200)
        row = response.json()[0]
        self.assertEqual(row['member_name'], 'Sara Benali')
        self.assertEqual(row['id_number'], 'CIN-CASH')
        self.assertEqual(row['receipt_number'], f'FO-{self.payment.id:06d}')
        self.assertEqual(Decimal(str(row['amount'])), Decimal('150.00'))
        self.assertEqual(Decimal(str(row['remaining_balance'])), Decimal('250.00'))

    def test_payment_list_can_search_by_cin_and_receipt(self):
        by_cin = self.client.get('/api/fitness/payments?q=CIN-CASH')
        self.assertEqual(by_cin.status_code, 200)
        self.assertEqual(len(by_cin.json()), 1)
        by_receipt = self.client.get(f'/api/fitness/payments?q=FO-{self.payment.id:06d}')
        self.assertEqual(by_receipt.status_code, 200)
        self.assertEqual(len(by_receipt.json()), 1)
        missing = self.client.get('/api/fitness/payments?q=nobody-here')
        self.assertEqual(missing.json(), [])

    def test_record_payment_returns_receipt_fields(self):
        response = self.client.post(
            f'/api/fitness/memberships/{self.membership.id}/payments',
            data=json.dumps({'amount': '50.00', 'received_by': 'Admin', 'notes': 'Cash'}),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body['member_name'], 'Sara Benali')
        self.assertTrue(body['receipt_number'].startswith('FO-'))
        self.assertEqual(Decimal(str(body['remaining_balance'])), Decimal('200.00'))

    def test_payment_uses_logged_in_staff_name_instead_of_admin(self):
        reception = User.objects.create_user(
            username='sara.desk',
            first_name='Sara',
            last_name='Desk',
            is_staff=True,
        )
        Group.objects.get_or_create(name='Reception')[0].user_set.add(reception)
        self.client.force_login(reception)
        response = self.client.post(
            f'/api/fitness/memberships/{self.membership.id}/payments',
            data=json.dumps({'amount': '50.00', 'received_by': 'Admin', 'notes': 'Cash'}),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['received_by'], 'Sara Desk')

    def test_payment_keeps_an_explicit_received_by_name(self):
        response = self.client.post(
            f'/api/fitness/memberships/{self.membership.id}/payments',
            data=json.dumps({'amount': '50.00', 'received_by': 'Karim Restore', 'notes': 'Cash'}),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['received_by'], 'Karim Restore')

    def test_receipt_html_and_pdf_are_available(self):
        training_class = TrainingClass.objects.create(name='Boxing Team', class_type=FitnessClassType.BOXING)
        ClassMember.objects.create(training_class=training_class, client=self.member)
        html = self.client.get(f'/api/fitness/payments/{self.payment.id}/receipt.html')
        self.assertEqual(html.status_code, 200)
        self.assertIn('text/html', html['Content-Type'])
        self.assertContains(html, 'Sara')
        self.assertContains(html, 'Benali')
        self.assertContains(html, 'Boxing Team')
        self.assertContains(html, '150.00')
        self.assertContains(html, 'Admin')
        self.assertNotContains(html, 'CIN-CASH')
        self.assertNotContains(html, '0612345678')
        self.assertNotContains(html, 'Still owes')
        pdf = self.client.get(f'/api/fitness/payments/{self.payment.id}/receipt')
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf['Content-Type'], 'application/pdf')
        self.assertTrue(pdf.content.startswith(b'%PDF'))

    def test_reception_can_export_cash_log(self):
        from io import BytesIO
        from openpyxl import load_workbook

        reception = User.objects.create_user(username='cash-reception')
        Group.objects.get_or_create(name='Reception')[0].user_set.add(reception)
        self.client.force_login(reception)
        url = f'/api/fitness/payments/export/xlsx?year={self.today.year}&month={self.today.month}'
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            f'AUMB-cash-log-{self.today.year}-{self.today.month:02d}.xlsx',
            response['Content-Disposition'],
        )
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.active['A5'].value, f'FO-{self.payment.id:06d}')
        self.assertEqual(workbook.active['C5'].value, 'Sara Benali')
        self.assertEqual(workbook.active['E5'].value, 150.0)
        pdf = self.client.get(f'/api/fitness/payments/export/pdf?year={self.today.year}&month={self.today.month}')
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.content.startswith(b'%PDF'))


class ThousandMemberQueryTests(TestCase):
    """Local scale check: 1000 members should not run one SQL query per row."""

    def setUp(self):
        from datetime import timedelta
        from decimal import Decimal

        from django.utils import timezone

        from fitness.models import GymPayment, Membership, MembershipPlan
        from users.models import ClientProfile

        staff = User.objects.create_user(username='scale-staff', password='password123', is_staff=True)
        self.client.force_login(staff)
        plan = MembershipPlan.objects.create(name='Scale plan', duration_months=1, price=Decimal('100.00'))
        today = timezone.localdate()
        User.objects.bulk_create([
            User(username=f'scale{i:04d}', first_name='Member', last_name=f'{i:04d}')
            for i in range(1000)
        ])
        users = list(User.objects.filter(username__startswith='scale').exclude(username='scale-staff').order_by('id'))
        ClientProfile.objects.bulk_create([
            ClientProfile(
                user=user,
                phone='0611223344',
                address='Rue Atlas',
                city='Casablanca',
                id_number=f'SC{i:06d}',
                is_active=True,
            )
            for i, user in enumerate(users)
        ])
        profiles = list(ClientProfile.objects.filter(id_number__startswith='SC').order_by('id'))
        Membership.objects.bulk_create([
            Membership(
                member=profile,
                plan=plan,
                start_date=today,
                end_date=today + timedelta(days=30),
                price=Decimal('100.00'),
            )
            for profile in profiles
        ])
        memberships = list(Membership.objects.order_by('id'))
        GymPayment.objects.bulk_create([
            GymPayment(
                membership=membership,
                amount=Decimal('100.00'),
                status='paid',
                payment_method='cash',
                received_by='desk',
            )
            for membership in memberships
        ])

    def test_member_and_membership_lists_stay_constant_query(self):
        from django.db import connection
        from django.test.utils import CaptureQueriesContext

        with CaptureQueriesContext(connection) as members_ctx:
            members = self.client.get('/api/fitness/members')
        with CaptureQueriesContext(connection) as memberships_ctx:
            memberships = self.client.get('/api/fitness/memberships')
        with CaptureQueriesContext(connection) as dash_ctx:
            dashboard = self.client.get('/api/fitness/dashboard')
        self.assertEqual(members.status_code, 200)
        self.assertEqual(len(members.json()), 500)
        self.assertLessEqual(len(members_ctx.captured_queries), 8)
        self.assertEqual(memberships.status_code, 200)
        self.assertEqual(len(memberships.json()), 500)
        self.assertLessEqual(len(memberships_ctx.captured_queries), 8)
        self.assertEqual(dashboard.status_code, 200)
        self.assertEqual(dashboard.json()['members'], 1000)
        self.assertLessEqual(len(dash_ctx.captured_queries), 16)


class Member360ApiTest(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(username='360-admin', password='password123', is_staff=True)
        self.today = timezone.localdate()
        self.user_a = User.objects.create_user(
            username='360-member-a',
            first_name='Sara',
            last_name='Benali',
            email='sara@example.com',
        )
        self.member_a = ClientProfile.objects.create(
            user=self.user_a,
            phone='0611111111',
            id_number='CIN360A',
            address='12 Rue Atlas',
            city='Casablanca',
        )
        self.user_b = User.objects.create_user(username='360-member-b', first_name='Lina', last_name='Said')
        self.member_b = ClientProfile.objects.create(user=self.user_b, phone='0622222222', id_number='CIN360B')
        self.plan = MembershipPlan.objects.create(name='Monthly', duration_months=1, price=Decimal('400.00'))
        self.older = Membership.objects.create(
            member=self.member_a,
            plan=self.plan,
            start_date=self.today - timedelta(days=60),
            end_date=self.today - timedelta(days=30),
            price=Decimal('400.00'),
            notes='Previous month',
        )
        self.current = Membership.objects.create(
            member=self.member_a,
            plan=self.plan,
            start_date=self.today - timedelta(days=1),
            end_date=self.today + timedelta(days=7),
            price=Decimal('400.00'),
            notes='Current plan',
        )
        self.other_membership = Membership.objects.create(
            member=self.member_b,
            plan=self.plan,
            start_date=self.today,
            end_date=self.today + timedelta(days=30),
            price=Decimal('400.00'),
            notes='Other member',
        )
        self.pay_a = GymPayment.objects.create(
            membership=self.current,
            amount=Decimal('100.00'),
            received_by='Desk',
            notes='A payment',
        )
        GymPayment.objects.create(
            membership=self.other_membership,
            amount=Decimal('50.00'),
            received_by='Desk',
            notes='B payment',
        )
        self.training_class = TrainingClass.objects.create(
            name='Morning Boxing',
            class_type=FitnessClassType.BOXING,
        )
        ClassMember.objects.create(training_class=self.training_class, client=self.member_a)
        Attendance.objects.create(member=self.member_a)
        Attendance.objects.create(member=self.member_b)
        GymWhatsAppReminder.objects.create(
            membership=self.current,
            kind='expiring_soon,unpaid',
            message='Due',
            sent_by='Desk',
        )

    def test_anonymous_gets_401(self):
        response = self.client.get(f'/api/fitness/members/{self.member_a.id}/360')
        self.assertEqual(response.status_code, 401)

    def test_gym_member_and_trainer_get_403(self):
        member_login = User.objects.create_user(username='360-self', password='password123')
        ClientProfile.objects.create(user=member_login, id_number='CIN360SELF')
        self.client.force_login(member_login)
        self.assertEqual(self.client.get(f'/api/fitness/members/{self.member_a.id}/360').status_code, 403)
        trainer = User.objects.create_user(username='360-trainer', password='password123')
        Group.objects.get_or_create(name='Trainer')[0].user_set.add(trainer)
        self.client.force_login(trainer)
        self.assertEqual(self.client.get(f'/api/fitness/members/{self.member_a.id}/360').status_code, 403)

    def test_reception_admin_and_superuser_can_read(self):
        reception = User.objects.create_user(username='360-reception', password='password123')
        Group.objects.get_or_create(name='Reception')[0].user_set.add(reception)
        self.client.force_login(reception)
        self.assertEqual(self.client.get(f'/api/fitness/members/{self.member_a.id}/360').status_code, 200)
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(f'/api/fitness/members/{self.member_a.id}/360').status_code, 200)
        superuser = User.objects.create_superuser(
            username='360-super',
            email='360-super@example.com',
            password='password123',
        )
        self.client.force_login(superuser)
        self.assertEqual(self.client.get(f'/api/fitness/members/{self.member_a.id}/360').status_code, 200)

    def test_missing_member_is_404(self):
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get('/api/fitness/members/999999/360').status_code, 404)

    def test_payload_is_scoped_to_the_member(self):
        self.client.force_login(self.admin)
        body = self.client.get(f'/api/fitness/members/{self.member_a.id}/360').json()
        member = body['member']
        self.assertEqual(member['id'], self.member_a.id)
        self.assertEqual(member['name'], 'Sara Benali')
        self.assertEqual(member['id_number'], 'CIN360A')
        self.assertEqual(member['phone'], '0611111111')
        self.assertTrue(member['is_active'])
        self.assertEqual(member['class_id'], self.training_class.id)
        self.assertEqual(body['training_class']['name'], 'Morning Boxing')
        membership_ids = [item['id'] for item in body['memberships']]
        self.assertEqual(membership_ids[0], self.current.id)
        self.assertEqual(set(membership_ids), {self.older.id, self.current.id})
        self.assertNotIn(self.other_membership.id, membership_ids)
        self.assertEqual(body['memberships'][0]['plan']['name'], 'Monthly')
        self.assertEqual(body['memberships'][0]['notes'], 'Current plan')
        payment_ids = [item['id'] for item in body['payments']]
        self.assertEqual(payment_ids, [self.pay_a.id])
        self.assertEqual(body['payments'][0]['notes'], 'A payment')
        attendance_members = {item['member_id'] for item in body['attendance']}
        self.assertEqual(attendance_members, {self.member_a.id})
        self.assertEqual(body['reminder']['member_id'], self.member_a.id)
        self.assertEqual(body['reminder']['membership_id'], self.current.id)

    def test_empty_history_is_200(self):
        empty_user = User.objects.create_user(username='360-empty', first_name='Empty')
        empty = ClientProfile.objects.create(user=empty_user, id_number='CIN360E')
        self.client.force_login(self.admin)
        body = self.client.get(f'/api/fitness/members/{empty.id}/360').json()
        self.assertEqual(body['member']['id'], empty.id)
        self.assertEqual(body['memberships'], [])
        self.assertEqual(body['payments'], [])
        self.assertEqual(body['attendance'], [])
        self.assertIsNone(body['reminder'])
        self.assertIsNone(body['training_class'])


class ClassScheduleCalendarApiTest(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(username='cal-admin', is_staff=True)
        self.reception = User.objects.create_user(username='cal-reception')
        Group.objects.get_or_create(name='Reception')[0].user_set.add(self.reception)
        self.boxing = TrainingClass.objects.create(name='Morning Boxing', class_type=FitnessClassType.BOXING)
        self.aerobic = TrainingClass.objects.create(name='Aerobic', class_type=FitnessClassType.AEROBIC)
        self.trainer = Trainer.objects.create(
            first_name='Karim',
            last_name='Coach',
            monthly_pay=Decimal('7777.00'),
        )
        member_user = User.objects.create_user(username='cal-roster', first_name='Sara')
        member = ClientProfile.objects.create(user=member_user, id_number='CAL-001')
        ClassMember.objects.create(training_class=self.boxing, client=member)
        self.monday = ClassSchedule.objects.create(
            training_class=self.boxing,
            weekday=0,
            start_time=time(18, 0),
            end_time=time(19, 0),
            trainer=self.trainer,
            location='Ring 1',
            capacity=12,
        )
        self.wednesday = ClassSchedule.objects.create(
            training_class=self.aerobic,
            weekday=2,
            start_time=time(10, 0),
            end_time=time(11, 0),
        )

    def _calendar(self, start='2026-09-01', end='2026-09-30'):
        return self.client.get(f'/api/fitness/classes/calendar?from={start}&to={end}')

    def test_anonymous_gets_401(self):
        self.assertEqual(self._calendar().status_code, 401)

    def test_gym_member_and_trainer_get_403(self):
        member_login = User.objects.create_user(username='cal-member', password='password123')
        ClientProfile.objects.create(user=member_login, id_number='CAL-MEM')
        self.client.force_login(member_login)
        self.assertEqual(self._calendar().status_code, 403)
        trainer_login = User.objects.create_user(username='cal-trainer-login', password='password123')
        Group.objects.get_or_create(name='Trainer')[0].user_set.add(trainer_login)
        self.client.force_login(trainer_login)
        self.assertEqual(self._calendar().status_code, 403)

    def test_reception_and_admin_can_read_calendar(self):
        self.client.force_login(self.reception)
        reception = self._calendar()
        self.assertEqual(reception.status_code, 200)
        self.assertEqual(self.client.get('/api/fitness/classes/schedules').status_code, 200)
        self.client.force_login(self.admin)
        admin = self._calendar()
        self.assertEqual(admin.status_code, 200)
        self.assertEqual(admin.json()['timezone'], 'Africa/Casablanca')
        super_admin = User.objects.create_user(username='cal-super')
        Group.objects.get_or_create(name='Super Admin')[0].user_set.add(super_admin)
        self.client.force_login(super_admin)
        self.assertEqual(self._calendar().status_code, 200)

    def test_reception_cannot_mutate_schedules(self):
        self.client.force_login(self.reception)
        payload = json.dumps({
            'training_class_id': self.boxing.id,
            'weekday': 'friday',
            'start_time': '09:00:00',
            'end_time': '10:00:00',
        })
        self.assertEqual(
            self.client.post(
                '/api/fitness/classes/schedules',
                data=payload,
                content_type='application/json',
            ).status_code,
            403,
        )
        self.assertEqual(
            self.client.put(
                f'/api/fitness/classes/schedules/{self.monday.id}',
                data=payload,
                content_type='application/json',
            ).status_code,
            403,
        )
        self.assertEqual(
            self.client.delete(f'/api/fitness/classes/schedules/{self.monday.id}').status_code,
            403,
        )
        self.assertTrue(ClassSchedule.objects.filter(id=self.monday.id).exists())

    def test_admin_can_create_update_and_deactivate_schedule(self):
        self.client.force_login(self.admin)
        created = self.client.post(
            '/api/fitness/classes/schedules',
            data=json.dumps({
                'training_class_id': self.boxing.id,
                'weekday': 'friday',
                'start_time': '09:00:00',
                'end_time': '10:00:00',
                'location': 'Studio',
                'group': 'kids',
                'capacity': 8,
                'is_active': True,
            }),
            content_type='application/json',
        )
        self.assertEqual(created.status_code, 200)
        body = created.json()
        self.assertEqual(body['weekday'], 'friday')
        self.assertEqual(body['location'], 'Studio')
        self.assertEqual(body['group'], 'kids')
        self.assertEqual(body['capacity'], 8)
        self.assertIsNone(body['trainer_id'])
        schedule_id = body['id']

        updated = self.client.put(
            f'/api/fitness/classes/schedules/{schedule_id}',
            data=json.dumps({
                'training_class_id': self.boxing.id,
                'weekday': 'friday',
                'start_time': '09:00:00',
                'end_time': '10:00:00',
                'trainer_id': self.trainer.id,
                'location': 'Studio 2',
                'group': 'adults',
                'capacity': 8,
                'is_active': False,
            }),
            content_type='application/json',
        )
        self.assertEqual(updated.status_code, 200)
        self.assertFalse(updated.json()['is_active'])
        self.assertEqual(updated.json()['group'], 'adults')
        self.assertEqual(updated.json()['trainer_name'], 'Karim Coach')
        self.assertNotIn('monthly_pay', updated.json())
        self.assertEqual(
            self._calendar(start='2026-09-04', end='2026-09-04').json()['items'],
            [],
        )
        super_admin = User.objects.create_user(username='cal-super-write')
        Group.objects.get_or_create(name='Super Admin')[0].user_set.add(super_admin)
        self.client.force_login(super_admin)
        deleted = self.client.delete(f'/api/fitness/classes/schedules/{schedule_id}')
        self.assertEqual(deleted.status_code, 200)
        self.assertFalse(ClassSchedule.objects.filter(id=schedule_id).exists())

    def test_weekly_occurrences_are_generated_inside_the_requested_range(self):
        self.client.force_login(self.admin)
        body = self._calendar().json()
        self.assertEqual(body['start_date'], '2026-09-01')
        self.assertEqual(body['end_date'], '2026-09-30')
        self.assertEqual(body['timezone'], 'Africa/Casablanca')
        dates = [item['date'] for item in body['items']]
        self.assertEqual(
            dates,
            [
                '2026-09-02', '2026-09-07', '2026-09-09', '2026-09-14',
                '2026-09-16', '2026-09-21', '2026-09-23', '2026-09-28',
                '2026-09-30',
            ],
        )
        monday = next(item for item in body['items'] if item['date'] == '2026-09-07')
        self.assertEqual(monday['schedule_id'], self.monday.id)
        self.assertEqual(monday['training_class_id'], self.boxing.id)
        self.assertEqual(monday['class_name'], 'Morning Boxing')
        self.assertEqual(monday['class_type'], 'boxing')
        self.assertEqual(monday['weekday'], 'monday')
        self.assertTrue(monday['start_time'].startswith('18:00'))
        self.assertTrue(monday['end_time'].startswith('19:00'))
        self.assertTrue(monday['starts_at'].startswith('2026-09-07T18:00:00'))
        self.assertIn('+01:00', monday['starts_at'])
        self.assertEqual(monday['trainer_id'], self.trainer.id)
        self.assertEqual(monday['trainer_name'], 'Karim Coach')
        self.assertEqual(monday['location'], 'Ring 1')
        self.assertEqual(monday['group'], '')
        self.assertEqual(monday['capacity'], 12)
        self.assertEqual(monday['member_count'], 1)
        self.assertNotIn('available_spaces', monday)
        encoded = json.dumps(body)
        self.assertNotIn('monthly_pay', encoded)
        self.assertNotIn('7777', encoded)
        self.assertNotIn('pay_amount', encoded)

    def test_schedule_outside_range_is_excluded(self):
        self.client.force_login(self.admin)
        body = self._calendar(start='2026-09-07', end='2026-09-07').json()
        self.assertEqual([item['date'] for item in body['items']], ['2026-09-07'])
        self.assertEqual(body['items'][0]['class_name'], 'Morning Boxing')
        self.assertEqual(len(body['items']), 1)

    def test_inactive_schedule_and_inactive_class_produce_no_occurrences(self):
        self.client.force_login(self.admin)
        self.monday.is_active = False
        self.monday.save(update_fields=['is_active', 'updated_at'])
        self.wednesday.training_class.is_active = False
        self.wednesday.training_class.save(update_fields=['is_active', 'updated_at'])
        body = self._calendar().json()
        self.assertEqual(body['items'], [])
        self.assertTrue(ClassSchedule.objects.filter(id=self.monday.id).exists())

    def test_invalid_time_and_date_ranges_are_rejected(self):
        self.client.force_login(self.admin)
        invalid_time = self.client.post(
            '/api/fitness/classes/schedules',
            data=json.dumps({
                'training_class_id': self.boxing.id,
                'weekday': 'monday',
                'start_time': '19:00:00',
                'end_time': '18:00:00',
            }),
            content_type='application/json',
        )
        self.assertEqual(invalid_time.status_code, 400)
        self.assertEqual(self.client.get('/api/fitness/classes/calendar').status_code, 400)
        self.assertEqual(self._calendar(start='2026-09-30', end='2026-09-01').status_code, 400)
        self.assertEqual(self._calendar(start='not-a-date', end='2026-09-01').status_code, 400)
        self.assertEqual(self._calendar(start='2026-01-01', end='2026-03-31').status_code, 400)
        self.assertEqual(
            self.client.post(
                '/api/fitness/classes/schedules',
                data=json.dumps({
                    'training_class_id': self.boxing.id,
                    'weekday': 'funday',
                    'start_time': '18:00:00',
                    'end_time': '19:00:00',
                }),
                content_type='application/json',
            ).status_code,
            400,
        )
        self.assertEqual(
            self.client.post(
                '/api/fitness/classes/schedules',
                data=json.dumps({
                    'training_class_id': 999999,
                    'weekday': 'monday',
                    'start_time': '18:00:00',
                    'end_time': '19:00:00',
                }),
                content_type='application/json',
            ).status_code,
            404,
        )
        self.assertEqual(
            self.client.post(
                '/api/fitness/classes/schedules',
                data=json.dumps({
                    'training_class_id': self.boxing.id,
                    'weekday': 'monday',
                    'start_time': '18:00:00',
                    'end_time': '19:00:00',
                    'trainer_id': 999999,
                }),
                content_type='application/json',
            ).status_code,
            404,
        )

    def test_existing_class_list_is_unchanged(self):
        self.client.force_login(self.admin)
        listed = self.client.get('/api/fitness/classes')
        self.assertEqual(listed.status_code, 200)
        row = listed.json()[0]
        self.assertEqual(
            set(row),
            {'id', 'name', 'class_type', 'price_per_member', 'member_count', 'team_total', 'is_active'},
        )
        self.assertNotIn('weekday', row)
        self.assertNotIn('start_time', row)

    def test_anonymous_schedule_endpoints_get_401(self):
        payload = json.dumps({
            'training_class_id': self.boxing.id,
            'weekday': 'monday',
            'start_time': '18:00:00',
            'end_time': '19:00:00',
        })
        self.assertEqual(self.client.get('/api/fitness/classes/schedules').status_code, 401)
        self.assertEqual(self.client.get(f'/api/fitness/classes/schedules/{self.monday.id}').status_code, 401)
        self.assertEqual(
            self.client.post(
                '/api/fitness/classes/schedules',
                data=payload,
                content_type='application/json',
            ).status_code,
            401,
        )

    def test_reception_can_read_schedule_but_ids_cannot_bypass_writes(self):
        self.client.force_login(self.reception)
        detail = self.client.get(f'/api/fitness/classes/schedules/{self.monday.id}')
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(detail.json()['id'], self.monday.id)
        self.assertNotIn('monthly_pay', detail.json())
        self.assertEqual(self.client.get('/api/fitness/classes/schedules/999999').status_code, 404)
        self.assertEqual(
            self.client.put(
                f'/api/fitness/classes/schedules/{self.monday.id}',
                data=json.dumps({
                    'training_class_id': self.boxing.id,
                    'weekday': 'sunday',
                    'start_time': '08:00:00',
                    'end_time': '09:00:00',
                    'is_active': False,
                }),
                content_type='application/json',
            ).status_code,
            403,
        )
        self.monday.refresh_from_db()
        self.assertEqual(self.monday.weekday, 0)
        self.assertTrue(self.monday.is_active)

    def test_calendar_payload_excludes_admin_financial_data(self):
        self.client.force_login(self.admin)
        body = self._calendar().json()
        encoded = json.dumps(body)
        forbidden = (
            'monthly_pay', 'pay_amount', 'is_paid', 'remaining_balance', 'total_paid',
            'price_per_member', 'team_total', 'expense', 'notification', 'receipt',
        )
        for key in forbidden:
            self.assertNotIn(key, encoded)
        item = next(row for row in body['items'] if row['schedule_id'] == self.monday.id)
        self.assertEqual(
            set(item),
            {
                'schedule_id', 'training_class_id', 'class_name', 'class_type', 'date',
                'weekday', 'start_time', 'end_time', 'starts_at', 'ends_at',
                'trainer_id', 'trainer_name', 'location', 'group', 'color', 'capacity', 'member_count',
                'is_active',
            },
        )
        self.assertEqual(item['capacity'], 12)
        self.assertEqual(item['member_count'], 1)
        self.assertNotIn('available_spaces', item)

    def test_schedule_group_is_the_calendar_name(self):
        self.client.force_login(self.admin)
        created = self.client.post(
            '/api/fitness/classes/schedules',
            data=json.dumps({
                'training_class_id': self.boxing.id,
                'weekday': 'thursday',
                'start_time': '16:00:00',
                'end_time': '17:00:00',
                'group': 'boxing-kids',
                'color': '#ef735c',
            }),
            content_type='application/json',
        )
        self.assertEqual(created.status_code, 200)
        self.assertEqual(created.json()['group'], 'boxing-kids')
        self.assertEqual(created.json()['color'], '#ef735c')
        body = self._calendar(start='2026-09-03', end='2026-09-03').json()
        self.assertEqual(len(body['items']), 1)
        self.assertEqual(body['items'][0]['class_name'], 'Morning Boxing')
        self.assertEqual(body['items'][0]['group'], 'boxing-kids')
        self.assertEqual(body['items'][0]['color'], '#ef735c')
        invalid_color = self.client.post(
            '/api/fitness/classes/schedules',
            data=json.dumps({
                'training_class_id': self.boxing.id,
                'weekday': 'thursday',
                'start_time': '17:00:00',
                'end_time': '18:00:00',
                'color': 'red',
            }),
            content_type='application/json',
        )
        self.assertEqual(invalid_color.status_code, 400)
        too_long = self.client.post(
            '/api/fitness/classes/schedules',
            data=json.dumps({
                'training_class_id': self.boxing.id,
                'weekday': 'thursday',
                'start_time': '17:00:00',
                'end_time': '18:00:00',
                'group': 'x' * 81,
            }),
            content_type='application/json',
        )
        self.assertEqual(too_long.status_code, 400)

    def test_calendar_range_is_capped_at_62_days(self):
        self.client.force_login(self.admin)
        allowed = self._calendar(start='2026-01-01', end='2026-03-03')
        self.assertEqual(allowed.status_code, 200)
        self.assertEqual(allowed.json()['start_date'], '2026-01-01')
        self.assertEqual(allowed.json()['end_date'], '2026-03-03')
        self.assertTrue(all('2026-01-01' <= item['date'] <= '2026-03-03' for item in allowed.json()['items']))
        self.assertEqual(self._calendar(start='2026-01-01', end='2026-03-04').status_code, 400)
        self.assertEqual(self._calendar(start='', end='2026-09-01').status_code, 400)

    def test_year_boundary_week_stays_inside_requested_range(self):
        self.client.force_login(self.admin)
        body = self._calendar(start='2026-12-28', end='2027-01-03').json()
        dates = [item['date'] for item in body['items']]
        self.assertEqual(dates, ['2026-12-28', '2026-12-30'])
        self.assertTrue(all('2026-12-28' <= item['date'] <= '2027-01-03' for item in body['items']))
        self.assertNotIn('2027-01-04', dates)

    def test_leap_year_february_does_not_leak_into_march(self):
        self.client.force_login(self.admin)
        body = self._calendar(start='2024-02-01', end='2024-02-29').json()
        dates = [item['date'] for item in body['items']]
        self.assertEqual(
            dates,
            [
                '2024-02-05', '2024-02-07', '2024-02-12', '2024-02-14',
                '2024-02-19', '2024-02-21', '2024-02-26', '2024-02-28',
            ],
        )
        self.assertTrue(all(item['date'].startswith('2024-02-') for item in body['items']))

    def test_monday_to_sunday_week_includes_only_that_week(self):
        self.client.force_login(self.admin)
        body = self._calendar(start='2026-09-07', end='2026-09-13').json()
        dates = [item['date'] for item in body['items']]
        self.assertEqual(dates, ['2026-09-07', '2026-09-09'])
        self.assertEqual(body['items'][0]['weekday'], 'monday')
        self.assertEqual(body['items'][1]['weekday'], 'wednesday')
        self.assertTrue(body['items'][0]['starts_at'].endswith('+01:00'))
        self.assertTrue(body['items'][0]['ends_at'].endswith('+01:00'))

    def test_deleting_class_cascades_schedules_and_keeps_member_profiles(self):
        self.client.force_login(self.admin)
        member = ClientProfile.objects.get(id_number='CAL-001')
        schedule_id = self.monday.id
        deleted = self.client.delete(f'/api/fitness/classes/{self.boxing.id}')
        self.assertEqual(deleted.status_code, 200)
        self.assertFalse(TrainingClass.objects.filter(id=self.boxing.id).exists())
        self.assertFalse(ClassSchedule.objects.filter(id=schedule_id).exists())
        self.assertTrue(ClientProfile.objects.filter(id=member.id).exists())
        self.assertFalse(ClassMember.objects.filter(client=member, training_class_id=self.boxing.id).exists())


class MemberQrLookupApiTest(TestCase):
    def setUp(self):
        from bookings.tests.helpers import make_provider

        self.admin = User.objects.create_user(username='qr-admin', is_staff=True)
        self.reception = User.objects.create_user(username='qr-reception')
        Group.objects.get_or_create(name='Reception')[0].user_set.add(self.reception)
        self.super_admin = User.objects.create_user(username='qr-super')
        Group.objects.get_or_create(name='Super Admin')[0].user_set.add(self.super_admin)
        self.provider = make_provider(username='qr-provider')
        member_user = User.objects.create_user(
            username='qr-member-a',
            first_name='Sara',
            last_name='Benali',
            email='sara@example.com',
        )
        self.member_a = ClientProfile.objects.create(
            user=member_user,
            phone='0611111111',
            id_number='QR-CIN-A',
            address='Hay Mohammadi',
            city='Casablanca',
        )
        other_user = User.objects.create_user(username='qr-member-b', first_name='Omar', last_name='Ali')
        self.member_b = ClientProfile.objects.create(user=other_user, phone='0622222222', id_number='QR-CIN-B')

    def _lookup(self, token=None):
        value = self.member_a.qr_token if token is None else token
        return self.client.get(f'/api/fitness/members/qr/{value}')

    def test_anonymous_gets_401(self):
        self.assertEqual(self._lookup().status_code, 401)

    def test_gym_member_trainer_and_provider_get_403(self):
        member_login = User.objects.create_user(username='qr-self', password='password123')
        ClientProfile.objects.create(user=member_login, id_number='QR-SELF')
        self.client.force_login(member_login)
        self.assertEqual(self._lookup().status_code, 403)
        trainer = User.objects.create_user(username='qr-trainer', password='password123')
        Group.objects.get_or_create(name='Trainer')[0].user_set.add(trainer)
        self.client.force_login(trainer)
        self.assertEqual(self._lookup().status_code, 403)
        self.client.force_login(self.provider.user)
        self.assertEqual(self._lookup().status_code, 403)

    def test_reception_admin_super_admin_and_staff_can_resolve(self):
        for user in (self.reception, self.admin, self.super_admin):
            self.client.force_login(user)
            response = self._lookup()
            self.assertEqual(response.status_code, 200, user.username)
            body = response.json()
            self.assertEqual(body, {
                'member_id': self.member_a.id,
                'name': 'Sara Benali',
                'is_active': True,
            })

    def test_invalid_and_unknown_tokens_are_404(self):
        self.client.force_login(self.admin)
        self.assertEqual(self._lookup('not-a-real-token').status_code, 404)
        self.assertEqual(self._lookup('FO-%06d' % self.member_a.id).status_code, 404)
        self.assertEqual(self.client.get('/api/fitness/members/qr/').status_code, 404)

    def test_member_a_token_never_returns_member_b(self):
        self.client.force_login(self.admin)
        body = self._lookup(self.member_a.qr_token).json()
        self.assertEqual(body['member_id'], self.member_a.id)
        self.assertNotEqual(body['member_id'], self.member_b.id)
        other = self._lookup(self.member_b.qr_token).json()
        self.assertEqual(other['member_id'], self.member_b.id)
        self.assertNotEqual(self.member_a.qr_token, self.member_b.qr_token)
        self.assertFalse(self.member_a.qr_token.startswith('FO-'))
        self.assertNotEqual(self.member_a.qr_token, str(self.member_a.id))

    def test_response_does_not_expose_private_data(self):
        self.client.force_login(self.admin)
        body = self._lookup().json()
        self.assertEqual(set(body), {'member_id', 'name', 'is_active'})
        encoded = json.dumps(body)
        for leaked in (
            'QR-CIN-A', '0611111111', 'sara@example.com', 'Hay Mohammadi',
            'payments', 'memberships', 'attendance', 'monthly_pay',
            'expenses', 'notifications', 'id_number', 'phone', 'email',
            'qr_token', self.member_a.qr_token,
        ):
            self.assertNotIn(leaked, encoded)

    def test_inactive_member_still_resolves(self):
        self.member_a.is_active = False
        self.member_a.save(update_fields=['is_active', 'updated_at'])
        self.client.force_login(self.reception)
        body = self._lookup().json()
        self.assertEqual(body['member_id'], self.member_a.id)
        self.assertFalse(body['is_active'])
        self.assertEqual(Attendance.objects.filter(member=self.member_a).count(), 0)

    def test_existing_member_360_and_member_list_are_unchanged(self):
        self.client.force_login(self.admin)
        detail = self.client.get(f'/api/fitness/members/{self.member_a.id}')
        self.assertEqual(detail.status_code, 200)
        self.assertNotIn('qr_token', detail.json())
        self.assertEqual(detail.json()['card_code'], f'FO-{self.member_a.id:06d}')
        view_360 = self.client.get(f'/api/fitness/members/{self.member_a.id}/360')
        self.assertEqual(view_360.status_code, 200)
        self.assertEqual(view_360.json()['member']['id'], self.member_a.id)
        self.assertNotIn('qr_token', view_360.json()['member'])
        image = self.client.get(f'/api/fitness/members/{self.member_a.id}/qr')
        self.assertEqual(image.status_code, 200)
        self.assertIn('image/svg', image['Content-Type'])



